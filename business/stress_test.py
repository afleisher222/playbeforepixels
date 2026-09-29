#!/usr/bin/env python3
"""
Play Before Pixels (AlphaPlay LLC) - financial stress test.

Re-runnable companion to business/STRESS-TEST.md. It reads every fee, price,
mix weight and cost line it can from the financial model workbook
(business/PlayBeforePixels_Financial_Model.xlsx, input cells only, so the
workbook does not need to be recalculated first), then runs:

  1. unit economics per launch product per channel;
  2. one-at-a-time sensitivity (tornado) on 12-month revenue and profit;
  3. a 10,000-run Monte Carlo for Oct 2026 - Sep 2027 and calendar 2026;
  4. the arithmetic behind a $1,000,000-by-end-of-2026 goal;
  5. break-even (fixed monthly costs vs contribution per order);
  6. the levers that most raise P50 (median) profit.

Everything prints as Markdown to stdout. Nothing is written to disk unless
you pass --json PATH.

Usage
  python3 business/stress_test.py                       # defaults
  python3 business/stress_test.py --runs 20000 --seed 7
  python3 business/stress_test.py --inputs actuals.json # re-run with real data
  python3 business/stress_test.py --json /tmp/out.json  # also save numbers

Re-running with real data (after 60-90 days of sales)
  Put observed values in a JSON file and pass it with --inputs. Any key of
  BASE (point estimates, used by the sensitivity, break-even and $1M
  sections) or MC (Monte Carlo ranges) can be overridden, e.g.

  {
    "BASE": {"launch": 3, "cvr_etsy": 0.017, "vpp_etsy": 210,
             "cvr_site": 0.011, "vpp_site": 140, "items_site": 1.4,
             "repeat12": 0.08, "refund": 0.015},
    "MC":   {"traction_median": 1.0, "traction_sigma": 0.4,
             "launch_probs": {"3": 1.0}}
  }

  Once BASE holds observed rates, set MC traction_median to 1.0 (the centre
  is then your own data, not the plan's Expected) and narrow traction_sigma
  as the months of data grow. launch_probs keys are model months (1 = Oct
  2026, 2 = Nov 2026, 3 = Dec 2026, ...).

  Where each number comes from once the shop is live:
    vpp_*      monthly visits per live listing (Etsy Stats "visits",
               Shopify Analytics "sessions", Gumroad views) / listings live
    cvr_*      orders / visits on that channel
    pvpt_kdp   KDP does not report page views; back it out as
               units per title per month / cvr_kdp
    items_*    items per order (Shopify, Etsy order exports)
    repeat12   share of buyers who order again within 12 months
    refund     refunds / gross sales

Labels used in the output: SOURCE (<repo file>) for numbers taken from the
repository; ASSUMPTION for planning inputs chosen here; UNVERIFIED for
platform rules or prices that have not been checked on the live page.
"""

from __future__ import annotations

import argparse
import copy
import json
import math
import sys
from pathlib import Path

import numpy as np

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

HERE = Path(__file__).resolve().parent
DEFAULT_WB = HERE / "PlayBeforePixels_Financial_Model.xlsx"

T = 15  # months simulated: Oct 2026 (t=0) .. Dec 2027 (t=14)
MONTH_NAMES = ["Oct 2026", "Nov 2026", "Dec 2026", "Jan 2027", "Feb 2027", "Mar 2027",
               "Apr 2027", "May 2027", "Jun 2027", "Jul 2027", "Aug 2027", "Sep 2027",
               "Oct 2027", "Nov 2027", "Dec 2027"]
CAL_MONTH = [10, 11, 12, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12]
Y1 = slice(0, 12)     # Oct 2026 - Sep 2027 (model year 1)
CY26 = slice(0, 3)    # Oct - Dec 2026 (calendar 2026)

# --------------------------------------------------------------------------
# Fallback inputs: identical to the workbook's Assumptions / Unit Economics /
# Monthly Operating Costs / Startup Costs input cells on 2026-09-28. Used only
# if the workbook cannot be read.
# --------------------------------------------------------------------------
FALLBACK_A = {
    "shop_pct": 0.029, "shop_fix": 0.30, "etsy_tx": 0.065, "etsy_pay": 0.03,
    "etsy_payfix": 0.25, "etsy_list": 0.20, "etsy_oa": 0.15, "etsy_oashare": 0.10,
    "kdp_roy": 0.60, "kdp_base": 1.00, "kdp_color": 0.07, "kdp_bw_flat": 2.84, "kdp_returns": 0.05,
    "pic_pages": 32, "ing_disc": 0.40, "ing_pb_print": 3.24, "mor_pct": 0.10,
    "mor_fix": 0.50, "mor_card_pct": 0.029, "mor_card_fix": 0.30, "refund": 0.05, "refund_course": 0.05, "contin": 0.10,
    "pos": 0.0, "first_sale": 3, "ramp": 6,
    "season": [1.25, 0.9, 1.0, 1.05, 0.95, 1.0, 0.95, 0.85, 0.85, 0.95, 1.3, 1.35],
}
FALLBACK_UE = [  # (channel, product, price, mix weight)
    ("SITE", "Visual routine cards", 9.5, 25), ("SITE", "\"I'm bored\" play cards", 6.5, 15),
    ("SITE", "Play-First Family Kit", 11.0, 20), ("SITE", "Toddler busy book printable", 11.99, 18),
    ("SITE", "100 Screen-Free Plays (PDF)", 9.99, 10), ("SITE", "Car Ride & Waiting Pack", 6.0, 5),
    ("SITE", "First-words flash cards", 6.99, 3),
    ("KDP", "100 Screen-Free Plays (paperback, B/W)", 16.99, 50),
    ("KDP", "The Day the Tablet Slept (paperback)", 11.99, 20),
    ("KDP", "Up! Go! More! (paperback)", 11.99, 30),
    ("COURSE", "30 Days of Back-and-Forth", 27.0, 70), ("COURSE", "30 Days of Back-and-Forth bundle", 49.0, 30),
]
FALLBACK_OPEX = [  # item, category, low, high, freq, start, end, include
    ("Shopify Basic plan", "Store", 39, 39, "Monthly", 1, 36, 1),
    ("Business email on the brand domain", "Tools", 7, 14, "Monthly", 1, 36, 1),
    ("Password manager", "Tools", 0, 40, "Annual", 1, 36, 1),
    ("Email platform, year 1 (free tiers, then budget)", "Email", 0, 20, "Monthly", 1, 12, 1),
    ("Email platform, years 2-3", "Email", 20, 39, "Monthly", 13, 36, 1),
    ("Claude plan that runs the scheduled routines", "Automation", 100, 200, "Monthly", 1, 36, 1),
    ("Claude usage-credit cap (hard limit)", "Automation", 0, 25, "Monthly", 1, 36, 1),
    ("QuickBooks Online (Simple Start to Essentials)", "Bookkeeping", 38, 85, "Monthly", 1, 36, 1),
    ("Link My Books settlement connector (1-3 channels)", "Connectors", 21, 47, "Monthly", 3, 36, 1),
    ("Buyer-specific PDF stamping", "Tools", 0, 20, "Monthly", 3, 36, 1),
    ("Social scheduler", "Tools", 0, 30, "Monthly", 3, 36, 1),
    ("Own-site review app", "Tools", 0, 15, "Monthly", 3, 36, 1),
    ("Etsy listing renewals (about 30 listings every 4 months)", "Channels", 1.5, 3, "Monthly", 3, 36, 1),
    ("Sales-tax filing service (optional; the accountant may file)", "Tax", 0, 25, "Monthly", 3, 36, 1),
    ("Domains, MUST tier", "Domains", 61, 61, "Annual", 1, 36, 1),
    ("Domains, SHOULD tier (before paid ads / first print run)", "Domains", 118, 118, "Annual", 5, 36, 1),
    ("USPS PO Box (public business address)", "Admin", 100, 300, "Annual", 1, 36, 1),
    ("Commercial resident agent", "Admin", 50, 300, "Annual", 1, 36, 1),
    ("Maryland SDAT annual report", "Admin", 300, 300, "Annual", 7, 36, 1),
    ("GDPR Art. 27 representatives, EU and UK (international digital sales)", "Legal/privacy", 220, 1300, "Annual", 4, 36, 1),
    ("General liability with products-completed operations", "Insurance", 45.17, 125, "Monthly", 3, 36, 1),
    ("Professional liability / E&O (switch: off until the broker advises)", "Insurance", 62, 125, "Monthly", 3, 36, 0),
    ("Media liability / publisher's E&O", "Insurance", 50, 150, "Monthly", 3, 36, 1),
    ("Cyber", "Insurance", 35, 129, "Monthly", 3, 36, 1),
    ("Umbrella / excess over GL", "Insurance", 25, 50, "Monthly", 5, 36, 1),
    ("Trademark watch (optional; the monthly DIY search is $0)", "Legal/IP", 0, 750, "Annual", 4, 36, 1),
    ("Accountant: year-end return and review", "Finance", 500, 1500, "Annual", 7, 36, 1),
    ("Copyright registrations for new releases (about 4 a year)", "Legal/IP", 260, 340, "Annual", 13, 36, 1),
    ("Upload assistant for KDP/IngramSpark/TpT packets", "Contractor", 0, 200, "Monthly", 7, 36, 0),
]
FALLBACK_STARTUP = [  # low, high, month, include  (21 included items, lean total $4,941)
    (25, 50, 1, 1), (25, 50, 1, 1), (6, 6, 1, 1), (0, 15, 1, 1), (295, 295, 1, 1),
    (500, 2500, 1, 1), (700, 700, 2, 1), (150, 625, 5, 1), (300, 1000, 4, 1),
    (360, 680, 2, 1), (500, 2000, 3, 1), (500, 2500, 1, 1), (300, 1000, 5, 1),
    (0, 300, 1, 1), (500, 1500, 6, 1), (300, 1000, 1, 1), (30, 90, 2, 1),
    (200, 500, 3, 1), (150, 600, 3, 1), (100, 300, 3, 1), (0, 20, 2, 1),
]


def _num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v)
        except ValueError:
            return None
    return None


def read_workbook(path: Path) -> dict:
    """Read input cells from the workbook. Returns dict with A, UE, OPEX, STARTUP, source."""
    out = {"A": dict(FALLBACK_A), "UE": list(FALLBACK_UE), "OPEX": list(FALLBACK_OPEX),
           "STARTUP": list(FALLBACK_STARTUP), "source": "fallback constants in stress_test.py",
           # workbook Dashboard year-1 gross sales, Expected and Strong (LibreOffice recalculation, Sep 29, 2026)
           "PLAN": {"exp": 11945.6, "strong": 29042.8}}
    if openpyxl is None or not path.exists():
        return out
    try:  # cached values exist only after a recalculation; otherwise keep the fallback
        wsd = openpyxl.load_workbook(path, data_only=True)["Dashboard"]
        for rr in range(1, wsd.max_row + 1):
            if wsd.cell(rr, 1).value == "Gross sales, year 1":
                e_, s_ = _num(wsd.cell(rr, 3).value), _num(wsd.cell(rr, 4).value)
                if e_ and s_:
                    out["PLAN"] = {"exp": e_, "strong": s_}
    except Exception:
        pass
    wb = openpyxl.load_workbook(path, data_only=False)
    A = dict(FALLBACK_A)
    ws = wb["Assumptions"]
    season = list(FALLBACK_A["season"])
    for r in range(1, ws.max_row + 1):
        key = ws.cell(r, 1).value
        val = _num(ws.cell(r, 3).value)
        if not isinstance(key, str) or val is None:
            continue
        if key.startswith("season_"):
            season[int(key.split("_")[1]) - 1] = val
        else:
            A[key] = val
    A["season"] = season
    ue = []
    ws = wb["Unit Economics"]
    for r in range(6, ws.max_row + 1):
        code, prod, price, mix = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 5).value, ws.cell(r, 16).value
        if code in ("SITE", "KDP", "COURSE") and _num(price) is not None and _num(mix) is not None:
            ue.append((code, str(prod), _num(price), _num(mix)))
    opex = []
    ws = wb["Monthly Operating Costs"]
    for r in range(6, ws.max_row + 1):
        item = ws.cell(r, 1).value
        low, high = _num(ws.cell(r, 3).value), _num(ws.cell(r, 4).value)
        freq, start, end = ws.cell(r, 6).value, _num(ws.cell(r, 7).value), _num(ws.cell(r, 8).value)
        inc = _num(ws.cell(r, 14).value)
        if item and low is not None and high is not None and freq in ("Monthly", "Annual") and start:
            opex.append((str(item), str(ws.cell(r, 2).value), low, high, freq, int(start), int(end), int(inc or 0)))
    startup = []
    ws = wb["Startup Costs"]
    for r in range(6, ws.max_row + 1):
        low, high = _num(ws.cell(r, 3).value), _num(ws.cell(r, 4).value)
        month, inc = _num(ws.cell(r, 6).value), _num(ws.cell(r, 7).value)
        if low is not None and high is not None and month is not None and inc is not None:
            startup.append((low, high, int(month), int(inc)))
    if ue:
        out["UE"] = ue
    if opex:
        out["OPEX"] = opex
    if startup:
        out["STARTUP"] = startup
    out["A"] = A
    out["source"] = str(path.relative_to(HERE.parent)) if HERE.parent in path.parents else str(path)
    return out


# --------------------------------------------------------------------------
# Base (point-estimate) inputs. Calibrated to the workbook's Expected case so
# the tornado starts from the plan's own measuring stick.
# --------------------------------------------------------------------------
BASE = {
    # timing
    "launch": 3,            # model month of first sale; 3 = Dec 2026. SOURCE Assumptions first_sale
    "ramp": 6,              # months to full visibility. SOURCE Assumptions ramp (ASSUMPTION there)
    # traffic at full ramp, per live listing (or title) per month, normal season
    "traffic_mult": 1.0,
    "vpp_etsy": 350.0,      # = 7 orders/product (SOURCE ETSY_cvr Expected) / 2.0% conversion (ASSUMPTION)
    "vpp_site": 266.7,      # = 4 orders/product (SOURCE SITE_cvr) / 1.5% (SOURCE site_cvr, MARKETING-PLAYBOOK)
    "vpp_mor": 80.0,        # = 1.2 orders/product (SOURCE MOR_cvr) / 1.5% (ASSUMPTION)
    "pvpt_kdp": 100.0,      # = 5 units/title (SOURCE KDP_cvr) / 5% page->order (ASSUMPTION)
    "upt_ing": 1.0,         # IngramSpark units/title/month (SOURCE INGRAM_cvr Expected)
    "noise_etsy": 1.0, "noise_site": 1.0, "noise_mor": 1.0, "noise_kdp": 1.0,
    # conversion
    "cvr_mult": 1.0,
    "cvr_etsy": 0.020, "cvr_site": 0.015, "cvr_mor": 0.015, "cvr_kdp": 0.05,
    # order value
    "aov_mult": 1.0,        # multiplies the price per digital order (bundles, order bumps)
    "items_site": 1.25, "items_etsy": 1.20, "items_mor": 1.20,   # SOURCE Assumptions *_items Expected
    # repeat purchases
    "repeat12": 0.10,       # share of own-site / MoR buyers who buy again within 12 months (ASSUMPTION)
    "etsy_repeat_factor": 0.5,  # Etsy buyers cannot be emailed (SOURCE sections/02 2.6), so half the rate
    # refunds and fees
    "refund": 0.05, "refund_course": 0.05,   # SOURCE Assumptions refund, refund_course (listing.json nets use 5%)
    "fee_mult": 1.0,        # multiplies Etsy/Shopify/Gumroad fees and Amazon's share of KDP list price
    "fee_set": 0.0,         # 0 = workbook fees; 1 = corrected (large-trim KDP print, Gumroad card fee, Ingram access fee)
    "etsy_oashare": 0.10,   # SOURCE Assumptions etsy_oashare
    # catalog
    "prod_start": 5.0, "prod_add": 0.75, "prod_cap": 15.0, "growth": 0.005,   # SOURCE Assumptions 6a/6b Expected
    "kdp_base": 3.0, "kdp_add": 0.33, "kdp_cap": 10.0,
    "ing_base": 2.0, "ing_add": 0.15, "ing_cap": 5.0,
    # email list and course
    "signup": 0.03, "optin": 0.15, "churn": 0.015, "reach": 0.15, "course_cvr": 0.015,
    # costs
    "cost_pos": 0.0,        # 0 = lean path (SOURCE Assumptions pos)
    "fixed_cut": 0.0,       # $ per month removed from the fixed run-rate (lever test)
    # paid ads (none in the base: the task specifies no ad budget)
    "ads_monthly": 0.0, "ads_cpc": 0.75, "ads_cvr": 0.07,   # Amazon Ads on KDP titles; CPC and cvr ASSUMPTION/UNVERIFIED
}

# Monte Carlo ranges for a new faceless Etsy/Shopify/KDP brand with no ad budget.
MC = {
    "launch_probs": {"1": 0.05, "2": 0.25, "3": 0.40, "4": 0.20, "5": 0.10},  # Oct..Feb first sale
    "traction_median": 0.60,   # common traffic factor vs the workbook's Expected; ~Conservative
    "traction_sigma": 0.80,    # lognormal sigma: P10 ~0.22x, P90 ~1.67x of Expected
    "channel_sigma": 0.35,     # independent channel noise
    "cvr_sigma": 0.30,         # lognormal sigma on each channel's conversion
    "aov_sigma": 0.12,
    "repeat12": (0.03, 0.10, 0.25),        # triangular
    "refund": (0.02, 0.05, 0.08),          # centred on the listings' 5% allowance (ASSUMPTION)
    "fee_mult": (0.95, 1.00, 1.20),
    "ramp": (4, 6, 10),
    "prod_add": (0.40, 0.75, 1.00),
    "prod_cap": (10, 15, 20),
    "growth": (0.0, 0.01),                  # uniform
    "cost_pos": (0.0, 0.1, 0.5),
    "fee_set_p_corrected": 1.0,             # use the corrected fee set in every run
}


# --------------------------------------------------------------------------
# Unit economics
# --------------------------------------------------------------------------
def fee_constants(A: dict, fee_set):
    """Channel fee constants. fee_set may be a float or numpy array (0..1)."""
    # fee_set 0 = exactly the workbook. fee_set 1 = "corrected": the repo's own later findings
    #   - KDP 8x10 and 8.5x8.5 are large trim: $2.84 flat B/W <=108 pp; $1.00 + $0.08/pp premium colour
    #     (business/REVENUE-PLAN.md rank 4 and 6 + critic; UNVERIFIED)
    #   - Gumroad adds 2.9% + $0.30 card processing to 10% + $0.50 (REVENUE-PLAN rank 1, marked VERIFIED there)
    #   - IngramSpark 1.875% market-access fee and colour print about $3.50-$4.50 (REVENUE-PLAN rank 6; UNVERIFIED)
    #   - the course sells through the merchant of record, not Shopify
    #     (products/course-screen-reset/listing.json channels; operations/AUTOMATION-MAP.md 3G)
    # Since Sep 29, 2026 the workbook itself carries the large-trim B/W print ($2.84), the Gumroad card fee and the
    # course on the merchant of record (all from listing.json), so the corrected set only adds what is still missing.
    fs = fee_set
    kdp_bw = A["kdp_bw_flat"] + fs * max(0.0, 2.84 - A["kdp_bw_flat"])
    kdp_col = (A["kdp_base"] + A["kdp_color"] * A["pic_pages"]) * (1 - fs) + fs * (1.00 + 0.08 * A["pic_pages"])
    card_pct, card_fix = A.get("mor_card_pct", 0.0), A.get("mor_card_fix", 0.0)
    mor_pct = A["mor_pct"] + card_pct + fs * max(0.0, 0.029 - card_pct)
    mor_fix = A["mor_fix"] + card_fix + fs * max(0.0, 0.30 - card_fix)
    return {
        "site_pct": A["shop_pct"], "site_fix": A["shop_fix"],
        "etsy_pct_base": A["etsy_tx"] + A["etsy_pay"], "etsy_oa": A["etsy_oa"],
        "etsy_fix": A["etsy_list"] + A["etsy_payfix"],
        "mor_pct": mor_pct, "mor_fix": mor_fix,
        "kdp_roy": A["kdp_roy"], "kdp_bw": kdp_bw, "kdp_col": kdp_col, "kdp_ret": A.get("kdp_returns", 0.0),
        "ing_disc": A["ing_disc"], "ing_fee": fs * 0.01875,
        "ing_bw": kdp_bw, "ing_col": A["ing_pb_print"] + fs * (4.00 - A["ing_pb_print"]),
        "course_pct": mor_pct,   # the course sells through the merchant of record (course-screen-reset/listing.json)
        "course_fix": mor_fix,
    }


def short(name):
    if "(" in name and not any(k in name for k in ("PDF", "paperback")):
        return name.split(" (")[0]
    return name


def blends(UE):
    def blend(code):
        rows = [(p, w, n) for c, n, p, w in UE if c == code and w > 0]
        W = sum(w for _, w, _ in rows)
        return rows, W
    return {c: blend(c) for c in ("SITE", "KDP", "COURSE")}


def net_per_item(price, pct, fix, refund, fm):
    return price * (1 - pct * fm - refund) - fix * fm


def unit_economics_table(A, UE):
    """Rows for the report: launch products x channel, workbook fees vs corrected."""
    f0, f1 = fee_constants(A, 0.0), fee_constants(A, 1.0)
    rows = []
    site = [(n, p) for c, n, p, w in UE if c == "SITE" and w > 0]
    for name, p in site:
        for ch in ("Own site (Shopify)", "Etsy", "Gumroad (MoR, outside US)"):
            if ch.startswith("Gumroad") and ("Car Ride" in name or "flash" in name):
                continue
            if ch.startswith("Own"):
                pct, fix, pct1, fix1 = f0["site_pct"], f0["site_fix"], f1["site_pct"], f1["site_fix"]
                label = "2.9% + $0.30 card (UNVERIFIED)"
            elif ch == "Etsy":
                pct = f0["etsy_pct_base"] + f0["etsy_oa"] * A["etsy_oashare"]
                fix = f0["etsy_fix"]
                pct1, fix1 = pct, fix
                label = "6.5% + 3% + $0.25 + $0.20 listing + 15% Offsite Ads on 10% (UNVERIFIED)"
            else:
                pct, fix, pct1, fix1 = f0["mor_pct"], f0["mor_fix"], f1["mor_pct"], f1["mor_fix"]
                label = "10% + $0.50 + 2.9% + $0.30 card (listing.json nets; UNVERIFIED)"
            n0 = net_per_item(p, pct, fix, A["refund"], 1.0)
            n1 = net_per_item(p, pct1, fix1, A["refund"], 1.0)
            rows.append((short(name), ch, p, p - n0 - p * A["refund"], p * A["refund"], 0.0, n0, n0 / p, n1, label))
    for name, p in [(n, p) for c, n, p, w in UE if c == "KDP"]:
        bw = "100 Screen" in name
        title = name.split(" (")[0]
        pr0 = f0["kdp_bw"] if bw else f0["kdp_col"]
        pr1 = f1["kdp_bw"] if bw else f1["kdp_col"]
        n0 = (p * f0["kdp_roy"] - pr0) * (1 - f0["kdp_ret"])
        n1 = (p * f1["kdp_roy"] - pr1) * (1 - f1["kdp_ret"])
        rows.append((title + " paperback", "Amazon KDP", p, p * (1 - f0["kdp_roy"]), (p * f0["kdp_roy"] - pr0) * f0["kdp_ret"], pr0, n0, n0 / p, n1,
                     f"60% royalty at $9.99+ (UNVERIFIED), 5% returns; print ${pr0:.2f} workbook vs ${pr1:.2f} large trim (UNVERIFIED)"))
        if "Tablet" not in name:
            pi0 = f0["ing_bw"] if bw else f0["ing_col"]
            pi1 = f1["ing_bw"] if bw else f1["ing_col"]
            ni0 = p * (1 - f0["ing_disc"]) - pi0
            ni1 = p * (1 - f1["ing_disc"] - f1["ing_fee"]) - pi1
            rows.append((title + " paperback", "IngramSpark, 40% discount", p, p * f0["ing_disc"], 0.0, pi0, ni0,
                         ni0 / p, ni1, f"40% discount (ASSUMPTION); print ${pi0:.2f} workbook proxy vs ${pi1:.2f}; "
                                       "1.875% access fee in corrected (UNVERIFIED)"))
    for name, p in [(n, p) for c, n, p, w in UE if c == "COURSE"]:
        n0 = net_per_item(p, f0["course_pct"], f0["course_fix"], A["refund_course"], 1.0)
        n1 = net_per_item(p, f1["course_pct"], f1["course_fix"], A["refund_course"], 1.0)
        rows.append((name, "Gumroad (merchant of record)", p, p * f0["course_pct"] + f0["course_fix"],
                     p * A["refund_course"], 0.0, n0, n0 / p, n1,
                     "Gumroad 10% + $0.50 + 2.9% + $0.30 card (listing.json route; UNVERIFIED); 5% refunds (ASSUMPTION)"))
    return rows


def net_lookup(A, UE, code, needle, channel="site"):
    """Net per unit at workbook fees for one product, used by the break-even and ads tables."""
    f0 = fee_constants(A, 0.0)
    for c, n, p, w in UE:
        if c == code and needle in n:
            if code == "KDP":
                pr = f0["kdp_bw"] if "100 Screen" in n else f0["kdp_col"]
                return p, (p * f0["kdp_roy"] - pr) * (1 - f0["kdp_ret"])
            if code == "COURSE":
                return p, net_per_item(p, f0["course_pct"], f0["course_fix"], A["refund_course"], 1.0)
            if channel == "etsy":
                return p, net_per_item(p, f0["etsy_pct_base"] + f0["etsy_oa"] * A["etsy_oashare"], f0["etsy_fix"], A["refund"], 1.0)
            return p, net_per_item(p, f0["site_pct"], f0["site_fix"], A["refund"], 1.0)
    raise KeyError(needle)


# --------------------------------------------------------------------------
# Fixed costs
# --------------------------------------------------------------------------
LAUNCH_TIED_STARTS = {3: 0, 4: 1, 5: 2}   # workbook start month -> offset from the first-sale month


def opex_matrix(OPEX, launch, cost_pos, contin, fixed_cut):
    """Monthly operating costs, shape (N, T). Items tied to the first sale move with it."""
    launch = np.asarray(launch, dtype=float)
    N = launch.shape[0]
    m = np.arange(1, T + 1)[None, :].repeat(N, 0)
    tot = np.zeros((N, T))
    for item, cat, low, high, freq, start, end, inc in OPEX:
        if not inc:
            continue
        amt = low + np.asarray(cost_pos, dtype=float).reshape(-1, 1) * (high - low)
        if start in LAUNCH_TIED_STARTS and not item.startswith("Domains") and "GDPR" not in item:
            s = launch.reshape(-1, 1) + LAUNCH_TIED_STARTS[start]
        elif "GDPR" in item:
            s = launch.reshape(-1, 1) + 1       # needed when international sales open (launch + 1)
        else:
            s = np.full((N, 1), float(start))
        live = (m >= s) & (m <= end)
        if freq == "Monthly":
            tot += amt * live
        else:
            due = live & (((m - s) % 12) == 0)
            tot += amt * due
    tot *= (1 + contin)
    tot -= np.asarray(fixed_cut, dtype=float).reshape(-1, 1) * (m >= 1)
    return np.maximum(tot, 0)


def startup_vector(STARTUP, cost_pos):
    v = np.zeros(T)
    for low, high, month, inc in STARTUP:
        if inc and 1 <= month <= T:
            v[month - 1] += low + cost_pos * (high - low)
    return v


def steady_state_fixed(OPEX, cost_pos, contin, year=2):
    """Amortized monthly run-rate by item (annual items / 12) for a steady-state month."""
    rows = []
    for item, cat, low, high, freq, start, end, inc in OPEX:
        if not inc:
            continue
        if year == 2 and not (start <= 18 <= end):
            continue
        if year == 1 and not (start <= 9 <= end):
            continue
        amt = low + cost_pos * (high - low)
        monthly = amt if freq == "Monthly" else amt / 12.0
        rows.append((item, cat, monthly))
    sub = sum(r[2] for r in rows)
    return rows, sub, sub * contin


# --------------------------------------------------------------------------
# The simulation (vectorised over runs)
# --------------------------------------------------------------------------
def as_arr(P, N):
    return {k: (np.full(N, float(v)) if np.isscalar(v) else np.asarray(v, dtype=float)) for k, v in P.items()}


def simulate(P: dict, W: dict, keep_monthly=False):
    A, UE, OPEX, STARTUP = W["A"], W["UE"], W["OPEX"], W["STARTUP"]
    N = max((len(v) for v in P.values() if not np.isscalar(v)), default=1)
    p = as_arr(P, N)
    col = lambda k: p[k].reshape(-1, 1)

    m = np.arange(1, T + 1)[None, :].astype(float)            # model month #
    season_raw = np.array(A["season"])
    season = np.array([season_raw[c - 1] for c in CAL_MONTH]) / season_raw.mean()
    season = season[None, :]

    L = col("launch")
    R = np.maximum(col("ramp"), 1)

    def ramp_from(start):
        k = m - start + 1
        return np.clip(k / R, 0, 1) * (m >= start)

    since = np.maximum(m - L, 0)
    live = (m >= L)
    products = np.minimum(col("prod_cap"), col("prod_start") + col("prod_add") * since) * live
    index = (1 + col("growth")) ** since
    prod_months = products * index * season

    tm, cm = col("traffic_mult"), col("cvr_mult")
    # visitors (listing visits / sessions / page views) and new orders
    L_mor = L + 1
    L_ing = L + 2
    L_course = np.maximum(4, L + 1)
    vis_etsy = prod_months * col("vpp_etsy") * tm * col("noise_etsy") * ramp_from(L)
    vis_site = prod_months * col("vpp_site") * tm * col("noise_site") * ramp_from(L)
    vis_mor = prod_months * col("vpp_mor") * tm * col("noise_mor") * ramp_from(L_mor) * (m >= L_mor)
    titles = np.minimum(col("kdp_cap"), col("kdp_base") + col("kdp_add") * since) * live
    vis_kdp = titles * season * col("pvpt_kdp") * tm * col("noise_kdp") * ramp_from(L)
    since_i = np.maximum(m - L_ing, 0)
    titles_i = np.minimum(col("ing_cap"), col("ing_base") + col("ing_add") * since_i) * (m >= L_ing)
    units_ing = titles_i * season * col("upt_ing") * tm * ramp_from(L_ing)

    new_etsy = vis_etsy * col("cvr_etsy") * cm
    new_site = vis_site * col("cvr_site") * cm
    new_mor = vis_mor * col("cvr_mor") * cm
    units_kdp = vis_kdp * col("cvr_kdp") * cm

    # repeat purchases from cumulative first-time buyers
    h = 1 - (1 - np.clip(col("repeat12"), 0, 0.99)) ** (1 / 12)
    h_etsy = 1 - (1 - np.clip(col("repeat12") * col("etsy_repeat_factor"), 0, 0.99)) ** (1 / 12)

    def with_repeat(new, hz):
        cum = np.cumsum(new, axis=1)
        prev = np.concatenate([np.zeros((new.shape[0], 1)), cum[:, :-1]], axis=1)
        return new + hz * prev

    ord_etsy = with_repeat(new_etsy, h_etsy)
    ord_site = with_repeat(new_site, h)
    ord_mor = with_repeat(new_mor, h)

    # email list -> course
    subs = np.zeros((N, T))
    course = np.zeros((N, T))
    prev = np.zeros(N)
    ramp_c = ramp_from(L_course) * (m >= L_course)
    for t in range(T):
        reached = prev * p["reach"] * season[0, t]
        course[:, t] = reached * p["course_cvr"] * p["cvr_mult"] * ramp_c[:, t]
        new_subs = vis_site[:, t] * p["signup"] + p["optin"] * (ord_site[:, t] + ord_mor[:, t] + units_kdp[:, t] + units_ing[:, t])
        subs[:, t] = prev * (1 - p["churn"]) + new_subs
        prev = subs[:, t]

    # prices and nets
    F = fee_constants(A, col("fee_set"))
    fm, rf, aov = col("fee_mult"), col("refund"), col("aov_mult")
    B = blends(UE)
    site_rows, Ws = B["SITE"]
    p_site = sum(pr * w for pr, w, _ in site_rows) / Ws
    mor_rows = [(pr, w) for pr, w, n in site_rows if "Car Ride" not in n and "flash" not in n]
    p_mor = sum(pr * w for pr, w in mor_rows) / sum(w for _, w in mor_rows)
    price_site, price_etsy, price_mor = p_site * aov, p_site * aov, p_mor * aov
    etsy_pct = F["etsy_pct_base"] + F["etsy_oa"] * col("etsy_oashare")
    net_site_i = net_per_item(price_site, F["site_pct"], F["site_fix"], rf, fm)
    net_etsy_i = net_per_item(price_etsy, etsy_pct, F["etsy_fix"], rf, fm)
    net_mor_i = net_per_item(price_mor, F["mor_pct"], F["mor_fix"], rf, fm)
    kdp_rows, Wk = B["KDP"]
    p_kdp = sum(pr * w for pr, w, _ in kdp_rows) / Wk
    net_kdp = sum(w * (pr * (1 - fm * (1 - F["kdp_roy"])) - (F["kdp_bw"] if "100 Screen" in n else F["kdp_col"])) * (1 - F["kdp_ret"])
                  for pr, w, n in kdp_rows) / Wk
    # IngramSpark carries 100 Plays + Up! Go! More! at equal weight (workbook INGRAM rows, 50/50)
    ing_rows = [(pr, 1.0, n) for pr, w, n in kdp_rows if "Tablet" not in n]
    Wi = sum(w for _, w, _ in ing_rows)
    p_ing = sum(pr * w for pr, w, _ in ing_rows) / Wi
    net_ing = sum(w * (pr * (1 - F["ing_disc"] - F["ing_fee"]) - (F["ing_bw"] if "100 Screen" in n else F["ing_col"]))
                  for pr, w, n in ing_rows) / Wi
    c_rows, Wc = B["COURSE"]
    p_course = sum(pr * w for pr, w, _ in c_rows) / Wc
    net_course = net_per_item(p_course, F["course_pct"], F["course_fix"], col("refund_course"), fm)

    it_s, it_e, it_m = col("items_site"), col("items_etsy"), col("items_mor")
    rev = {
        "Etsy": ord_etsy * it_e * price_etsy,
        "Own site": ord_site * it_s * price_site,
        "Gumroad (intl)": ord_mor * it_m * price_mor,
        "KDP": units_kdp * p_kdp,
        "IngramSpark": units_ing * p_ing,
        "Course": course * p_course,
    }
    con = {
        "Etsy": ord_etsy * it_e * net_etsy_i,
        "Own site": ord_site * it_s * net_site_i,
        "Gumroad (intl)": ord_mor * it_m * net_mor_i,
        "KDP": units_kdp * net_kdp,
        "IngramSpark": units_ing * net_ing,
        "Course": course * net_course,
    }
    orders = ord_etsy + ord_site + ord_mor + units_kdp + units_ing + course
    visitors = vis_etsy + vis_site + vis_mor + vis_kdp

    # paid ads (Amazon Ads on KDP titles), from the month after launch
    ads_on = (m >= L + 1)
    spend = col("ads_monthly") * ads_on
    ad_units = np.where(col("ads_cpc") > 0, spend / col("ads_cpc"), 0) * col("ads_cvr")
    rev["Ads-driven KDP"] = ad_units * p_kdp
    con["Ads-driven KDP"] = ad_units * net_kdp
    orders = orders + ad_units

    revenue = sum(rev.values())
    contrib = sum(con.values())
    opex = opex_matrix(OPEX, p["launch"], p["cost_pos"], A["contin"], p["fixed_cut"])
    one_time = np.vstack([startup_vector(STARTUP, cp) for cp in p["cost_pos"]]) if N <= 64 else \
        startup_vector(STARTUP, 0.0)[None, :] + p["cost_pos"].reshape(-1, 1) * (
            startup_vector(STARTUP, 1.0) - startup_vector(STARTUP, 0.0))[None, :]
    op = contrib - opex - spend
    cum = np.cumsum(op - one_time, axis=1)
    units_pp = col("traffic_mult") * col("cvr_mult") * (
        col("vpp_etsy") * col("noise_etsy") * col("cvr_etsy") * col("items_etsy")
        + col("vpp_site") * col("noise_site") * col("cvr_site") * col("items_site")
        + col("vpp_mor") * col("noise_mor") * col("cvr_mor") * col("items_mor"))
    res = {
        "peak_loss": np.maximum(0, -cum.min(axis=1)),
        "op_m12": op[:, 11],
        "units_pp": units_pp[:, 0],
        "rev_y1": revenue[:, Y1].sum(1), "rev_cy26": revenue[:, CY26].sum(1),
        "contrib_y1": contrib[:, Y1].sum(1), "contrib_cy26": contrib[:, CY26].sum(1),
        "opex_y1": opex[:, Y1].sum(1), "opex_cy26": opex[:, CY26].sum(1),
        "ads_y1": spend[:, Y1].sum(1),
        "op_y1": op[:, Y1].sum(1), "op_cy26": op[:, CY26].sum(1),
        "net_y1": (op - one_time)[:, Y1].sum(1), "net_cy26": (op - one_time)[:, CY26].sum(1),
        "orders_y1": orders[:, Y1].sum(1), "orders_cy26": orders[:, CY26].sum(1),
        "visitors_y1": visitors[:, Y1].sum(1), "visitors_cy26": visitors[:, CY26].sum(1),
        "visitors_m12": visitors[:, 11], "orders_m12": orders[:, 11],
        "contrib_m12": contrib[:, 11], "rev_m12": revenue[:, 11],
        "one_time_y1": one_time[:, Y1].sum(1),
        "subs_m12": subs[:, 11],
    }
    for k in rev:
        res["rev_y1_" + k] = rev[k][:, Y1].sum(1)
        res["con_y1_" + k] = con[k][:, Y1].sum(1)
    if keep_monthly:
        res["monthly"] = {"revenue": revenue, "contrib": contrib, "opex": opex, "orders": orders,
                          "visitors": visitors, "op": op, "one_time": one_time,
                          "vis_etsy": vis_etsy, "vis_site": vis_site, "vis_mor": vis_mor, "vis_kdp": vis_kdp}
    res["unit"] = {"p_site": p_site, "p_mor": p_mor, "p_kdp": p_kdp, "p_course": p_course,
                   "net_site_i": net_site_i, "net_etsy_i": net_etsy_i, "net_mor_i": net_mor_i,
                   "net_kdp": net_kdp, "net_ing": net_ing, "net_course": net_course}
    return res


# --------------------------------------------------------------------------
# Monte Carlo draws
# --------------------------------------------------------------------------
def tri(rng, lo_mode_hi, n):
    lo, mode, hi = lo_mode_hi
    if hi == lo:
        return np.full(n, float(lo))
    return rng.triangular(lo, mode, hi, n)


def draw_mc(rng, n, base, mc):
    P = {k: v for k, v in base.items()}
    months = np.array([int(k) for k in mc["launch_probs"]])
    probs = np.array([mc["launch_probs"][k] for k in mc["launch_probs"]], dtype=float)
    P["launch"] = rng.choice(months, size=n, p=probs / probs.sum()).astype(float)
    z0 = rng.standard_normal(n)
    P["traffic_mult"] = mc["traction_median"] * np.exp(mc["traction_sigma"] * z0)
    for ch in ("etsy", "site", "mor", "kdp"):
        P["noise_" + ch] = np.exp(mc["channel_sigma"] * rng.standard_normal(n) - 0.5 * mc["channel_sigma"] ** 2)
        P["cvr_" + ch] = base["cvr_" + ch] * np.exp(mc["cvr_sigma"] * rng.standard_normal(n))
    P["aov_mult"] = np.exp(mc["aov_sigma"] * rng.standard_normal(n))
    P["repeat12"] = tri(rng, mc["repeat12"], n)
    P["refund"] = tri(rng, mc["refund"], n)
    P["fee_mult"] = tri(rng, mc["fee_mult"], n)
    P["ramp"] = tri(rng, mc["ramp"], n)
    P["prod_add"] = tri(rng, mc["prod_add"], n)
    P["prod_cap"] = tri(rng, mc["prod_cap"], n)
    P["growth"] = rng.uniform(mc["growth"][0], mc["growth"][1], n)
    P["cost_pos"] = tri(rng, mc["cost_pos"], n)
    P["fee_set"] = (rng.uniform(size=n) < mc["fee_set_p_corrected"]).astype(float)
    P["ads_monthly"] = np.zeros(n)
    return P


def pct(x, q):
    return float(np.percentile(x, q))


# --------------------------------------------------------------------------
# Formatting helpers
# --------------------------------------------------------------------------
def usd(x, dp=0):
    s = f"{abs(x):,.{dp}f}"
    return f"(${s})" if x < 0 else f"${s}"


def num(x, dp=0):
    return f"{x:,.{dp}f}"


def table(headers, rows):
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    out += ["| " + " | ".join(str(c) for c in r) + " |" for r in rows]
    return "\n".join(out)


def one(res, k):
    return float(res[k][0])


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook", default=str(DEFAULT_WB))
    ap.add_argument("--inputs", help="JSON file with BASE / MC overrides (real data)")
    ap.add_argument("--runs", type=int, default=10000)
    ap.add_argument("--seed", type=int, default=20260928)
    ap.add_argument("--json", help="also write the key numbers to this JSON path")
    a = ap.parse_args(argv)

    base, mc = copy.deepcopy(BASE), copy.deepcopy(MC)
    if a.inputs:
        ov = json.loads(Path(a.inputs).read_text())
        base.update(ov.get("BASE", {}))
        for k, v in ov.get("MC", {}).items():
            mc[k] = tuple(v) if isinstance(v, list) else v
    W = read_workbook(Path(a.workbook))
    A = W["A"]
    base["launch"] = base.get("launch", A.get("first_sale", 3))
    out = {}
    P = lambda s: print(s)

    P(f"# Stress test output\n\nWorkbook inputs read from: `{W['source']}`. Runs: {a.runs:,}. Seed: {a.seed}.\n")

    # ---------------- 1. Unit economics
    P("## 1. Unit economics per launch product per channel\n")
    ue = unit_economics_table(A, W["UE"])
    P(table(["Product", "Channel", "Price", "Platform + payment fees", "Refund allowance", "Print cost",
             "Net (workbook fees)", "Margin", "Net (corrected fees)", "Fee basis"],
            [(r[0], r[1], usd(r[2], 2), usd(r[3], 2), usd(r[4], 2), usd(r[5], 2), usd(r[6], 2),
              f"{r[7]*100:.0f}%", usd(r[8], 2), r[9]) for r in ue]))
    out["unit_economics"] = [dict(zip(["product", "channel", "price", "fees", "refund", "print", "net", "margin",
                                       "net_corrected", "basis"], r)) for r in ue]

    # ---------------- base case
    b = simulate(base, W, keep_monthly=True)
    mo = b["monthly"]
    lm_name = MONTH_NAMES[int(base["launch"]) - 1]
    P(f"\n## Base case (workbook Expected demand unless overridden, first sale {lm_name}, "
      f"cost position {base['cost_pos']}, ads ${base['ads_monthly']:.0f}/mo)\n")
    P(table(["Measure", "Oct 2026 - Sep 2027", "Calendar 2026"],
            [("Gross sales", usd(one(b, "rev_y1")), usd(one(b, "rev_cy26"))),
             ("Net contribution", usd(one(b, "contrib_y1")), usd(one(b, "contrib_cy26"))),
             ("Operating costs", usd(one(b, "opex_y1")), usd(one(b, "opex_cy26"))),
             ("Operating profit (before one-time costs)", usd(one(b, "op_y1")), usd(one(b, "op_cy26"))),
             ("One-time startup costs", usd(one(b, "one_time_y1")), usd(float(mo["one_time"][0, CY26].sum()))),
             ("Profit after one-time costs", usd(one(b, "net_y1")), usd(one(b, "net_cy26"))),
             ("Orders", num(one(b, "orders_y1")), num(one(b, "orders_cy26"))),
             ("Visitors (Etsy visits + site sessions + Gumroad + Amazon book-page views)", num(one(b, "visitors_y1")), num(one(b, "visitors_cy26")))]))
    P("\nMonthly base path:\n")
    P(table(["Month", "Visitors", "Orders", "Gross sales", "Contribution", "Operating costs", "Operating profit"],
            [(MONTH_NAMES[t], num(mo["visitors"][0, t]), num(mo["orders"][0, t], 1), usd(mo["revenue"][0, t]),
              usd(mo["contrib"][0, t]), usd(mo["opex"][0, t]), usd(mo["op"][0, t])) for t in range(12)]))
    ch_rows = [(k, usd(one(b, "rev_y1_" + k)), usd(one(b, "con_y1_" + k))) for k in
               ("Etsy", "Own site", "Gumroad (intl)", "KDP", "IngramSpark", "Course")]
    P("\nBy channel, Oct 2026 - Sep 2027:\n")
    P(table(["Channel", "Gross sales", "Contribution"], ch_rows))
    out["base"] = {k: one(b, k) for k in b if isinstance(b[k], np.ndarray)}
    _, _sub, _cont = steady_state_fixed(W["OPEX"], 0.0, A["contin"], year=2)
    be_full_lean = _sub + _cont                                   # lean fixed run-rate, annual bills spread
    be_cpo = float(b["contrib_m12"][0] / b["orders_m12"][0])      # contribution per order at month 12

    # ---------------- 2. Sensitivity / tornado
    P("\n## 2. Sensitivity: one input at a time, from the base case\n")
    sens = [
        ("Monthly visitors", "traffic_mult", 0.5, 1.5, "-50% / +50%"),
        ("Conversion rate (all channels)", "cvr_mult", 0.7, 1.3, "-30% / +30%"),
        ("Average order value (digital)", "aov_mult", 0.85, 1.15, "-15% / +15%"),
        ("Repeat-purchase rate (12-month)", "repeat12", 0.03, 0.25, "3% / 25% (base 10%)"),
        ("Refund rate", "refund", 0.08, 0.02, "8% / 2% (base 5%)"),
        ("Platform fees", "fee_mult", 1.2, 0.8, "+20% / -20%"),
        ("Paid ads (Amazon Ads, CPC $0.75, 7% click-to-sale)", "ads_monthly", 300.0, 0.0, "$300/mo / $0 (base $0)"),
        ("First-sale month", "launch", 5.0, 2.0, "Feb 2027 / Nov 2026 (base Dec 2026 unless overridden)"),
        ("Fixed-cost position", "cost_pos", 0.5, 0.0, "mid-point / lean (base lean)"),
        ("Months to full visibility (ramp)", "ramp", 10.0, 4.0, "10 / 4 (base 6)"),
        ("Fee set", "fee_set", 1.0, 0.0, "corrected / workbook"),
        ("New products per month", "prod_add", 0.4, 1.0, "0.4 / 1.0 (base 0.75)"),
    ]
    b_rev, b_op = one(b, "rev_y1"), one(b, "op_y1")
    srows = []
    for label, key, lo, hi, desc in sens:
        r_lo = simulate({**base, key: lo}, W)
        r_hi = simulate({**base, key: hi}, W)
        srows.append((label, desc, one(r_lo, "rev_y1"), one(r_hi, "rev_y1"), one(r_lo, "op_y1"), one(r_hi, "op_y1")))
    srows.sort(key=lambda r: -abs(r[5] - r[4]))
    P(f"Base: 12-month gross sales {usd(b_rev)}, operating profit {usd(b_op)}. Ranked by the swing in operating profit.\n")
    P(table(["Rank", "Input", "Low / high tested", "Revenue (low)", "Revenue (high)", "Profit (low)", "Profit (high)", "Profit swing"],
            [(i + 1, r[0], r[1], usd(r[2]), usd(r[3]), usd(r[4]), usd(r[5]), usd(abs(r[5] - r[4])))
             for i, r in enumerate(srows)]))
    out["tornado"] = [dict(zip(["input", "range", "rev_lo", "rev_hi", "op_lo", "op_hi"], r)) for r in srows]

    # elasticities
    P("\nElasticity (percent change in 12-month operating profit for a 1% change in the input, around the base):\n")
    el = []
    for label, key in (("Monthly visitors", "traffic_mult"), ("Conversion rate", "cvr_mult"),
                       ("Average order value", "aov_mult"), ("Platform fees", "fee_mult"),
                       ("Refund rate", "refund"), ("Repeat-purchase rate", "repeat12")):
        v0 = base[key]
        r1 = simulate({**base, key: v0 * 1.01}, W)
        el.append((label, (one(r1, "op_y1") - b_op) / abs(b_op) * 100, (one(r1, "rev_y1") - b_rev) / b_rev * 100))
    P(table(["Input", "Profit elasticity", "Revenue elasticity"], [(e[0], f"{e[1]:+.2f}", f"{e[2]:+.2f}") for e in el]))

    # grid visitors x conversion
    P("\n12-month operating profit, visitors x conversion (other inputs at base):\n")
    vm = [0.25, 0.5, 1.0, 1.5, 2.0]
    cmv = [0.5, 0.75, 1.0, 1.25, 1.5]
    grid = []
    for v in vm:
        row = [f"{v:.2f}x visitors"]
        for c in cmv:
            row.append(usd(one(simulate({**base, "traffic_mult": v, "cvr_mult": c}, W), "op_y1")))
        grid.append(row)
    P(table(["", *[f"{c:.2f}x conversion" for c in cmv]], grid))

    # ads break-even CPC
    P("\nHighest cost per click that breaks even on the first sale (contribution per order x click-to-order rate):\n")
    u = b["unit"]
    ads_rows = []
    _rp, _rn = net_lookup(A, W["UE"], "SITE", "routine", "etsy")
    _bp, _bn = net_lookup(A, W["UE"], "SITE", "busy book", "etsy")
    _pp, _pn = net_lookup(A, W["UE"], "KDP", "100 Screen")
    _cp, _cn = net_lookup(A, W["UE"], "KDP", "Tablet")
    for name, contrib_order, cv in (
            (f"Visual routine cards on Etsy (${_rp:.2f})", _rn, 0.02),
            (f"Busy book on Etsy (${_bp:.2f})", _bn, 0.02),
            ("Blended own-site order", float(u["net_site_i"][0, 0]) * base["items_site"], 0.015),
            ("Ages 1-5 Instant Gift Bundle on own site ($29)", 29 * (1 - A["shop_pct"] - A["refund"]) - A["shop_fix"], 0.015),
            ("100 Screen-Free Plays paperback, Amazon Ads", _pn, 0.07),
            ("32-page colour paperback, Amazon Ads", _cn, 0.07)):
        ads_rows.append((name, usd(contrib_order, 2), f"{cv*100:.1f}%", usd(contrib_order * cv, 2)))
    P(table(["Product", "Contribution per order", "Click-to-order (ASSUMPTION)", "Break-even CPC"], ads_rows))

    # ---------------- 3. Monte Carlo
    rng = np.random.default_rng(a.seed)
    D = draw_mc(rng, a.runs, base, mc)
    r = simulate(D, W, keep_monthly=True)
    P(f"\n## 3. Monte Carlo ({a.runs:,} runs, no ad budget)\n")
    rows = []
    for lab, k in (("Gross sales, Oct 2026 - Sep 2027", "rev_y1"), ("Operating profit, Oct 2026 - Sep 2027", "op_y1"),
                   ("Profit after one-time costs, Oct 2026 - Sep 2027", "net_y1"),
                   ("Orders, Oct 2026 - Sep 2027", "orders_y1"),
                   ("Gross sales, calendar 2026", "rev_cy26"), ("Operating profit, calendar 2026", "op_cy26"),
                   ("Profit after one-time costs, calendar 2026", "net_cy26"), ("Orders, calendar 2026", "orders_cy26"),
                   ("Visitors in Sep 2027 (month 12)", "visitors_m12"), ("Orders in Sep 2027", "orders_m12"),
                   ("Operating profit in Sep 2027 alone", "op_m12"),
                   ("Deepest cumulative loss through Dec 2027 (founder-capital proxy, before payout delays)", "peak_loss"),
                   ("Digital units per product per month at full ramp (plan anchors: floor 2.5, Low 10, Expected 15, Strong 30)", "units_pp")):
        if "units" in lab:
            f = lambda x: num(x, 1)
        elif "Orders" in lab or "Visitors" in lab:
            f = num
        else:
            f = usd
        rows.append((lab, f(pct(r[k], 10)), f(pct(r[k], 50)), f(pct(r[k], 90)), f(float(r[k].mean()))))
    P(table(["Measure", "P10", "P50", "P90", "Mean"], rows))
    P("\nMedian (P50) 12-month gross sales by channel: " + "; ".join(
        f"{k} {usd(pct(r['rev_y1_' + k], 50))}" for k in ("Etsy", "Own site", "Gumroad (intl)", "KDP", "IngramSpark", "Course")) + ".\n")
    # which uncertain inputs drive the spread (Spearman rank correlation with 12-month operating profit)
    def ranks(x):
        return np.argsort(np.argsort(x)).astype(float)
    ry = ranks(r["op_y1"])
    drivers = []
    for lab_, key in (("Traffic (common traction factor)", "traffic_mult"), ("Fixed-cost position", "cost_pos"),
                      ("First-sale month (later = worse)", "launch"), ("Etsy conversion", "cvr_etsy"),
                      ("Own-site conversion", "cvr_site"), ("Amazon book-page conversion", "cvr_kdp"),
                      ("Order value multiplier", "aov_mult"), ("Months to full visibility", "ramp"),
                      ("New products per month", "prod_add"), ("Etsy traffic (channel-specific)", "noise_etsy"),
                      ("Own-site traffic (channel-specific)", "noise_site"), ("Platform-fee multiplier", "fee_mult"),
                      ("Growth in sales per product", "growth"), ("Catalog cap", "prod_cap"),
                      ("Refund rate", "refund"), ("Repeat-purchase rate", "repeat12")):
        x = np.asarray(D[key], dtype=float)
        if np.ptp(x) == 0:
            continue
        drivers.append((lab_, float(np.corrcoef(ranks(x), ry)[0, 1])))
    drivers.sort(key=lambda d: -abs(d[1]))
    P("\nWhat drives the spread across the 10,000 futures (rank correlation with 12-month operating profit; "
      "+1 or -1 = decides everything, 0 = no effect):\n")
    P(table(["Rank", "Uncertain input", "Rank correlation"], [(i + 1, d[0], f"{d[1]:+.2f}") for i, d in enumerate(drivers)]))
    out["mc_drivers"] = drivers
    be_orders = be_full_lean / be_cpo
    probs = [
        ("12-month operating profit above $0", (r["op_y1"] > 0).mean()),
        ("12-month profit after one-time costs above $0", (r["net_y1"] > 0).mean()),
        (f"Sep 2027 orders at or above the break-even line ({be_orders:.0f} a month, lean)", (r["orders_m12"] >= be_orders).mean()),
        ("Sep 2027 operating profit above $0", (r["op_m12"] > 0).mean()),
        ("Deepest cumulative loss above the $12,000 household-money cap placeholder", (r["peak_loss"] > 12000).mean()),
        (f"12-month gross sales above the workbook's Expected ({usd(W['PLAN']['exp'])})", (r["rev_y1"] > W["PLAN"]["exp"]).mean()),
        (f"12-month gross sales above the workbook's Strong ({usd(W['PLAN']['strong'])})", (r["rev_y1"] > W["PLAN"]["strong"]).mean()),
        ("Digital units per product below the kill-rule floor (2.5 a month)", (r["units_pp"] < 2.5).mean()),
        ("Any sale in calendar 2026", (r["rev_cy26"] > 0).mean()),
        ("Calendar-2026 gross sales of $1,000,000 or more", (r["rev_cy26"] >= 1e6).mean()),
    ]
    P("\n" + table(["Probability", "Share of runs"], [(p_[0], f"{p_[1]*100:.1f}%") for p_ in probs]))
    # when does the median future reach the break-even pace and its first 100 orders?
    om = r["monthly"]["orders"]
    hit = om >= be_orders
    first_be = np.where(hit.any(1), hit.argmax(1), T)
    cum = np.cumsum(om, axis=1)
    first_100 = np.where((cum >= 100).any(1), (cum >= 100).argmax(1), T)
    def mname(q, arr):
        i = int(np.percentile(arr, q, method="lower"))
        return MONTH_NAMES[i] if i < T else "after Dec 2027"
    P("\nMilestones (month reached, across runs):\n")
    P(table(["Milestone", "P25 (faster futures)", "P50", "P75 (slower futures)", "Not reached by Dec 2027"],
            [(f"Monthly orders first reach the break-even pace ({be_orders:.0f})", mname(25, first_be), mname(50, first_be),
              mname(75, first_be), f"{(first_be >= T).mean()*100:.0f}%"),
             ("Cumulative orders pass 100", mname(25, first_100), mname(50, first_100), mname(75, first_100),
              f"{(first_100 >= T).mean()*100:.0f}%")]))
    out["milestones"] = {"first_be_p50": mname(50, first_be), "first_100_p50": mname(50, first_100),
                         "never_be": float((first_be >= T).mean())}
    out["mc"] = {k: {"p10": pct(r[k], 10), "p50": pct(r[k], 50), "p90": pct(r[k], 90), "mean": float(r[k].mean())}
                 for k in ("rev_y1", "op_y1", "net_y1", "orders_y1", "rev_cy26", "op_cy26", "net_cy26",
                           "orders_cy26", "visitors_m12", "orders_m12", "visitors_cy26", "op_m12", "peak_loss",
                           "units_pp")}
    out["mc_probs"] = {p_[0]: float(p_[1]) for p_ in probs}
    # by launch month
    lrows = []
    for lm, name in ((1, "Oct 2026"), (2, "Nov 2026"), (3, "Dec 2026"), (4, "Jan 2027"), (5, "Feb 2027")):
        sel = D["launch"] == lm
        if sel.sum() == 0:
            continue
        lrows.append((name, f"{sel.mean()*100:.0f}%", usd(pct(r["rev_cy26"][sel], 50)), usd(pct(r["rev_cy26"][sel], 90)),
                      usd(pct(r["rev_y1"][sel], 50)), usd(pct(r["op_y1"][sel], 50))))
    P("\nBy first-sale month:\n")
    P(table(["First sale", "Share of runs", "CY2026 sales P50", "CY2026 sales P90", "12-mo sales P50", "12-mo op. profit P50"], lrows))

    # ---------------- 4. $1M
    P("\n## 4. What $1,000,000 by December 31, 2026 would take\n")
    m_ = b["monthly"]
    aov_plan = float(m_["revenue"][0, 2] / m_["orders"][0, 2])
    conv_plan = float(m_["orders"][0, 11] / m_["visitors"][0, 11])
    windows = (("Selling from today, Sep 28 (not possible: no bank account, counsel or insurance yet)", 95),
               ("First sale Nov 1 (the hoped-for launch)", 61), ("First sale Dec 1 (the plan's month)", 31),
               ("For comparison: the whole year Oct 2026 - Sep 2027", 365))
    mrows = []
    for label, days in windows:
        for aov_ in ((aov_plan,) if days == 365 else (aov_plan, 29.0, 49.0)):
            orders_ = 1e6 / aov_
            vis_ = orders_ / conv_plan
            mrows.append((label, f"{days} days ({days/7:.1f} wk)", usd(aov_, 2), num(orders_), num(orders_ / days),
                          num(vis_), num(vis_ / days)))
    P(f"Blended order value in the plan's first month: {usd(aov_plan, 2)}. Blended visitor-to-order conversion at month 12: {conv_plan*100:.2f}%.\n")
    P(table(["Window", "Days", "Order value", "Orders needed", "Orders per day", "Visitors needed", "Visitors per day"], mrows))
    p90_26 = pct(r["rev_cy26"], 90)
    p90_y1 = pct(r["rev_y1"], 90)
    P(f"\nMonte Carlo P90 for calendar-2026 gross sales: {usd(p90_26)} ({num(pct(r['orders_cy26'], 90))} orders). "
      f"$1,000,000 is {num(1e6 / max(p90_26, 1))} times that. P90 for the full 12 months to Sep 2027: {usd(p90_y1)}; "
      f"$1,000,000 is {num(1e6 / p90_y1, 1)} times that. Highest calendar-2026 result in {a.runs:,} runs: {usd(float(r['rev_cy26'].max()))}.")
    out["million"] = {"aov_plan": aov_plan, "conv_plan": conv_plan, "p90_cy26": p90_26, "p90_y1": p90_y1,
                      "max_cy26": float(r["rev_cy26"].max())}

    # ---------------- 5. Break-even
    P("\n## 5. Break-even\n")
    for cp, lab in ((0.0, "lean path"), (0.5, "mid-point")):
        items, sub, cont = steady_state_fixed(W["OPEX"], cp, A["contin"], year=2)
        named = {"Shopify": 0.0, "Domains": 0.0, "Email platform": 0.0, "Insurance": 0.0, "PO Box": 0.0,
                 "Accounting (QuickBooks, Link My Books, accountant)": 0.0}
        other = []
        for it, cat, mly in items:
            if it.startswith("Shopify"):
                named["Shopify"] += mly
            elif cat == "Domains":
                named["Domains"] += mly
            elif cat == "Email":
                named["Email platform"] += mly
            elif cat == "Insurance":
                named["Insurance"] += mly
            elif "PO Box" in it:
                named["PO Box"] += mly
            elif cat in ("Bookkeeping", "Connectors", "Finance"):
                named["Accounting (QuickBooks, Link My Books, accountant)"] += mly
            else:
                other.append((it, mly))
        out.setdefault("breakeven", {})[lab] = {"named": named, "other": other, "subtotal": sub, "contingency": cont}
    lean = out["breakeven"]["lean path"]
    mid = out["breakeven"]["mid-point"]
    rows = [(k, usd(v, 2), usd(mid["named"][k], 2)) for k, v in lean["named"].items()]
    named_lean = sum(lean["named"].values())
    named_mid = sum(mid["named"].values())
    rows.append(("**Subtotal: the six named costs**", usd(named_lean, 2), usd(named_mid, 2)))
    other_lean = sum(v for _, v in lean["other"])
    other_mid = sum(v for _, v in mid["other"])
    rows.append(("Everything else (Claude plan, business email, resident agent, SDAT, GDPR reps, Etsy renewals, copyright filings)",
                 usd(other_lean, 2), usd(other_mid, 2)))
    rows.append(("10% contingency", usd(lean["contingency"], 2), usd(mid["contingency"], 2)))
    full_lean = lean["subtotal"] + lean["contingency"]
    full_mid = mid["subtotal"] + mid["contingency"]
    rows.append(("**All fixed costs, per month (annual bills spread over 12)**", usd(full_lean, 2), usd(full_mid, 2)))
    P(table(["Fixed cost (steady state, year 2)", "Lean path", "Mid-point"], rows))
    cpo = float(b["contrib_m12"][0] / b["orders_m12"][0])
    P(f"\nBlended contribution per order at month 12 (base mix): {usd(cpo, 2)}.\n")
    be_rows = []
    for lab_, amt in (("Six named costs, lean", named_lean), ("All fixed costs, lean", full_lean),
                      ("All fixed costs, lean + $150/mo ads", full_lean + 150),
                      ("Six named costs, mid-point", named_mid), ("All fixed costs, mid-point", full_mid)):
        be_rows.append((lab_, usd(amt, 0), num(amt / cpo), num(amt / cpo / 30.4, 1)))
    P(table(["Costs to cover", "Per month", "Orders per month", "Orders per day"], be_rows))
    single = []
    for lab_, code_, needle_, ch_ in (("30 Days of Back-and-Forth, Gumroad", "COURSE", "Back-and-Forth", "site"),
                                      ("Play-First Family Kit, own site", "SITE", "Family Kit", "site"),
                                      ("Toddler busy book, Etsy", "SITE", "busy book", "etsy"),
                                      ("Visual routine cards, Etsy", "SITE", "routine", "etsy"),
                                      ("Play-First Family Kit, Etsy", "SITE", "Family Kit", "etsy"),
                                      ("100 Screen-Free Plays paperback, KDP", "KDP", "100 Screen", "site"),
                                      ("\"I'm bored\" play cards, Etsy", "SITE", "bored", "etsy"),
                                      ("32-page colour paperback, KDP", "KDP", "Tablet", "site")):
        pr_, nt_ = net_lookup(A, W["UE"], code_, needle_, ch_)
        single.append((f"{lab_} (${pr_:.2f})", nt_))
    single.sort(key=lambda x: -x[1])
    P("\nIf only one product sold (all fixed costs, lean path):\n")
    P(table(["Product", "Net per unit (workbook)", "Units per month", "Per day"],
            [(n, usd(v, 2), num(full_lean / v), num(full_lean / v / 30.4, 1)) for n, v in single]))
    out["breakeven"]["cpo"] = cpo
    out["breakeven"]["full_lean"] = full_lean
    out["breakeven"]["full_mid"] = full_mid

    # ---------------- 6. Levers
    P("\n## 6. Levers: change in P50 (median) 12-month operating profit, same random draws\n")
    p50_base = pct(r["op_y1"], 50)
    p50_rev = pct(r["rev_y1"], 50)

    def lever(mod):
        D2 = {k: (v.copy() if isinstance(v, np.ndarray) else v) for k, v in D.items()}
        mod(D2)
        r2 = simulate(D2, W)
        return (pct(r2["op_y1"], 50), pct(r2["rev_y1"], 50), pct(r2["op_y1"], 10), pct(r2["op_y1"], 90),
                pct(r2["peak_loss"], 50), float((r2["peak_loss"] > 12000).mean()))

    levers = [
        ("Cut $100/month of fixed cost by quote-shopping (insurance, Claude plan tier, bookkeeping stack)",
         lambda d: d.__setitem__("fixed_cut", np.full(a.runs, 100.0))),
        ("Be selling by Nov 1 (Gate A met in October)", lambda d: d.__setitem__("launch", np.minimum(d["launch"], 2.0))),
        ("Raise digital order value 15% (the $29 and $45 bundles, one order bump)",
         lambda d: d.__setitem__("aov_mult", d["aov_mult"] * 1.15)),
        ("Raise conversion 20% (mockups, previews, Start Here page, first reviews)",
         lambda d: d.__setitem__("cvr_mult", np.full(a.runs, 1.2))),
        ("Raise traffic 20% (Pinterest + Etsy SEO cadence, backlinks)",
         lambda d: d.__setitem__("traffic_mult", d["traffic_mult"] * 1.2)),
        ("Add 1 new product a month instead of 0.4-1.0 (more listings, cap 20)",
         lambda d: (d.__setitem__("prod_add", np.full(a.runs, 1.0)), d.__setitem__("prod_cap", np.full(a.runs, 20.0)))),
        ("Double the repeat-purchase rate (post-purchase emails, next-age offers)",
         lambda d: d.__setitem__("repeat12", np.minimum(d["repeat12"] * 2, 0.6))),
        ("Hold every cost at the lean-path quote (cost position 0)", lambda d: d.__setitem__("cost_pos", np.zeros(a.runs))),
        ("Keep Etsy Offsite Ads off while under $10k/yr", lambda d: d.__setitem__("etsy_oashare", np.zeros(a.runs))),
        ("Run $150/month Amazon Ads from the month after launch", lambda d: d.__setitem__("ads_monthly", np.full(a.runs, 150.0))),
    ]
    lrows = []
    for lab_, fn in levers:
        p50, rev50, p10, p90, pl50, pcap = lever(fn)
        lrows.append((lab_, p50 - p50_base, rev50 - p50_rev, p10, p50, p90, pl50, pcap))
    lrows.sort(key=lambda x: -x[1])
    P(f"Baseline P50: 12-month operating profit {usd(p50_base)}, gross sales {usd(p50_rev)}.\n")
    P(f"Baseline P50 deepest cumulative loss {usd(pct(r['peak_loss'], 50))}; runs past the $12,000 cap placeholder {(r['peak_loss'] > 12000).mean()*100:.0f}%.\n")
    P(table(["Rank", "Lever (effect size is an ASSUMPTION)", "Change in P50 profit", "Change in P50 sales", "New P10 profit",
             "New P50 profit", "New P90 profit", "P50 deepest loss", "Runs past $12k cap"],
            [(i + 1, x[0], usd(x[1]), usd(x[2]), usd(x[3]), usd(x[4]), usd(x[5]), usd(x[6]), f"{x[7]*100:.0f}%")
             for i, x in enumerate(lrows)]))
    out["levers"] = [dict(zip(["lever", "d_p50_op", "d_p50_rev", "p10", "p50", "p90", "p50_peak_loss", "p_cap"], x))
                     for x in lrows]

    if a.json:
        def conv(o):
            if isinstance(o, dict):
                return {k: conv(v) for k, v in o.items()}
            if isinstance(o, (list, tuple)):
                return [conv(v) for v in o]
            if isinstance(o, (np.floating, np.integer)):
                return float(o)
            return o
        Path(a.json).write_text(json.dumps(conv(out), indent=2))
    return out


if __name__ == "__main__":
    main()
