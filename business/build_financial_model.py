#!/usr/bin/env python3
"""Build PlayBeforePixels_Financial_Model.xlsx (AlphaPlay LLC d/b/a Play Before Pixels).

Every number is either sourced to a repo file or labelled Assumption / [VERIFY].
All downstream cells are live Excel formulas.
"""
import datetime
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, BarChart, Reference, Series
from openpyxl.comments import Comment

OUT = sys.argv[1] if len(sys.argv) > 1 else "/home/user/playbeforepixels/business/PlayBeforePixels_Financial_Model.xlsx"

# ---------------------------------------------------------------- styles
AR = "Arial"
def fnt(color="000000", bold=False, size=10, italic=False):
    return Font(name=AR, color=color, bold=bold, size=size, italic=italic)
F_IN, F_LINK, F_CALC = fnt("0000FF"), fnt("008000"), fnt("000000")
F_NOTE = fnt("555555", italic=True, size=9)
HDR_FILL = PatternFill("solid", fgColor="1D2940")
HDR_FONT = Font(name=AR, bold=True, color="FFFFFF", size=10)
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
SEC_FILL = PatternFill("solid", fgColor="E3E7EF")
TOT_FILL = PatternFill("solid", fgColor="F3F4F7")
THIN = Side(style="thin", color="B7BDC9")
TOP = Border(top=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CUR0 = '$#,##0;($#,##0);"-"'
CUR2 = '$#,##0.00;($#,##0.00);"-"'
PCT1 = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
NUM0 = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM2 = '0.00'
DATEF = 'mmm-yy'

def style_font(v, explicit=None):
    if explicit is not None:
        return explicit
    if isinstance(v, str) and v.startswith("="):
        return F_LINK if "!" in v else F_CALC
    return F_CALC

def put(ws, ref, v, fmt=None, font=None, fill=None, bold=False, wrap=False, align=None):
    c = ws[ref]
    c.value = v
    f = style_font(v, font)
    if bold:
        f = Font(name=AR, color=f.color, bold=True, size=f.size, italic=f.italic)
    c.font = f
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = WRAP
    if align:
        c.alignment = align
    return c

def hdr(ws, row, c1, c2, labels=None, height=None):
    for i, c in enumerate(range(c1, c2 + 1)):
        cell = ws.cell(row=row, column=c)
        if labels is not None and i < len(labels):
            cell.value = labels[i]
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if c > 1 else "left")
    if height:
        ws.row_dimensions[row].height = height

def sec(ws, row, c1, c2, text):
    for c in range(c1, c2 + 1):
        ws.cell(row=row, column=c).fill = SEC_FILL
    cell = ws.cell(row=row, column=c1)
    cell.value = text
    cell.font = fnt(bold=True)

def title(ws, text, sub):
    put(ws, "A1", text, font=Font(name=AR, bold=True, size=14, color="1D2940"))
    put(ws, "A2", sub, font=F_NOTE)

def q(sheet):
    return f"'{sheet}'"

wb = Workbook()
SHEETS = ["Dashboard", "Assumptions", "Unit Economics", "Startup Costs", "Monthly Operating Costs",
          "36-Month Forecast", "Cash Flow", "Break-even", "Retail Readiness Costs"]
ws_ = {}
ws_["Dashboard"] = wb.active
ws_["Dashboard"].title = "Dashboard"
for n in SHEETS[1:]:
    ws_[n] = wb.create_sheet(n)
for ws in ws_.values():
    ws.sheet_view.showGridLines = False

SCN = ["Conservative", "Expected", "Strong"]
SCOL = ["C", "D", "E"]
SERIES_COLORS = ["2A78D6", "EB6834", "1BAF7A"]

# =====================================================================
# ASSUMPTIONS
# =====================================================================
wa = ws_["Assumptions"]
title(wa, "Assumptions: every input in the model",
      "Blue = editable input. Yellow fill = key scenario lever. Black = formula. Green = link to another tab. "
      "Status: Repo = figure found in a repository file (most are themselves marked UNVERIFIED there); "
      "Assumption = planning input chosen for this model; [VERIFY] = outside fact not in the repo, check before relying on it.")
wa.merge_cells("A2:H3")
wa["A2"].alignment = WRAP
wa.row_dimensions[2].height = 30
hdr(wa, 5, 1, 8, ["ID", "Input", "Value / Conservative", "Expected", "Strong", "Unit", "Source / basis", "Status"], 30)
A = {}
R = [6]

def a_sec(text):
    R[0] += 1
    sec(wa, R[0], 1, 8, text)
    R[0] += 1

def a_sub(labels):
    hdr(wa, R[0], 1, 8, labels)
    R[0] += 1

def inp(key, label, vals, fmt, unit, src, status, key_lever=False, formula=False):
    r = R[0]
    put(wa, f"A{r}", key, font=F_NOTE)
    put(wa, f"B{r}", label, wrap=True)
    if isinstance(vals, (list, tuple)):
        A[key] = []
        for i, v in enumerate(vals):
            cref = f"{SCOL[i]}{r}"
            put(wa, cref, v, fmt=fmt, font=None if formula else F_IN, fill=KEY_FILL if key_lever else None)
            A[key].append(f"{q('Assumptions')}!${SCOL[i]}${r}")
    else:
        put(wa, f"C{r}", vals, fmt=fmt, font=None if formula else F_IN, fill=KEY_FILL if key_lever else None)
        A[key] = f"{q('Assumptions')}!$C${r}"
    put(wa, f"F{r}", unit)
    put(wa, f"G{r}", src, wrap=True)
    put(wa, f"H{r}", status)
    R[0] += 1
    return r

a_sec("1. General settings")
inp("start", "Model start month (month 1)", datetime.date(2026, 10, 1), DATEF, "month",
    "Plan dated Sept 28, 2026. Wave 1 runs Oct-Dec 2026 (marketing/BLIND-SPOTS.md #2).", "Repo")
inp("pos", "Cost position inside every Low-High range (0 = low, 1 = high)", 0.5, NUM2, "0-1",
    "Drives the 'Model' column on Startup Costs, Monthly Operating Costs and Retail Readiness Costs.", "Assumption", True)
inp("ramp", "Months for a new listing to reach full conversion (review ramp)", 6, "0", "months",
    "A new faceless shop starts with zero reviews against incumbents with 2,000-11,000 (marketing/DEMAND-CHECK.md, 'One honest warning').", "Assumption")
inp("mdtax", "Maryland sales tax on direct (own-site) sales", 0.06, PCT1, "% of price",
    "commerce/storefront-setup-guide.md A3 (UNVERIFIED). Collected on top of price and remitted: pass-through, never counted as revenue.", "Repo [VERIFY]")
inp("refund", "Refund and chargeback allowance, direct digital sales", 0.02, PCT1, "% of price",
    "ops/GAPS-ROUND-2.md G2-12: first refund request under about $15 is refunded automatically.", "Assumption")
inp("refund_course", "Refund allowance, 30-Day Screen Reset (money-back guarantee)", 0.05, PCT1, "% of price",
    "DEMAND-CHECK course row: 'Add a money-back guarantee'.", "Assumption")
inp("taxres", "Tax reserve transfer (% of each positive month's net cash)", 0.25, PCT1, "%",
    "finance/TAX-AUTOPILOT.md §2: accountant picks the rate; 25-30% is common for self-employment income.", "Repo range")
inp("open", "Opening cash in the business account", 0, CUR0, "$",
    "finance/BANKING.md: the Chase checking account is closed; a new no-fee account opens empty.", "Repo")
inp("owner", "Owner capital contribution in month 1", 0, CUR0, "$",
    "Left at $0 so the model shows the funding need. Enter what the founder will put in.", "Assumption", True)
inp("contin", "Operating contingency (% of monthly operating costs)", 0.10, PCT1, "%",
    "Buffer for price changes; nearly every subscription price in the repo is UNVERIFIED.", "Assumption")
inp("resmo", "Cash reserve target (months of fixed costs)", 3, "0", "months",
    "ops/ROUTINE.md weekly scorecard ('3-month reserve target'); BLIND-SPOTS #20 says 2-3 months.", "Repo")

a_sec("2. Platform fees and royalties (used by Unit Economics)")
inp("shop_pct", "Shopify Payments card rate", 0.029, PCT1, "% of price", "commerce/storefront-setup-guide.md §1; finance/money-and-tax-setup.md fees table (Basic plan, UNVERIFIED).", "Repo [VERIFY]")
inp("shop_fix", "Shopify Payments fixed fee", 0.30, CUR2, "$ per order", "Same source.", "Repo [VERIFY]")
inp("etsy_tx", "Etsy transaction fee", 0.065, PCT1, "% of price", "storefront-setup-guide §10 (UNVERIFIED).", "Repo [VERIFY]")
inp("etsy_pay", "Etsy payment processing (US)", 0.03, PCT1, "% of price", "Same source.", "Repo [VERIFY]")
inp("etsy_payfix", "Etsy payment processing fixed fee", 0.25, CUR2, "$ per order", "Same source.", "Repo [VERIFY]")
inp("etsy_list", "Etsy listing fee (renews on each sale)", 0.20, CUR2, "$ per sale", "Same source.", "Repo [VERIFY]")
inp("etsy_oa", "Etsy Offsite Ads fee when an ad drives the sale", 0.15, PCT1, "% of price", "Same source: 15%, or 12% and mandatory above $10,000 in 12 months.", "Repo [VERIFY]")
inp("etsy_oashare", "Share of Etsy sales that come through Offsite Ads", 0.10, PCT1, "% of Etsy sales", "No data yet; check the Etsy statement after 90 days.", "Assumption")
inp("kdp_roy", "KDP paperback royalty rate (list price $9.99 or more)", 0.60, PCT1, "% of list", "storefront-setup-guide §6: 60% at $9.99+ since June 10, 2025 (UNVERIFIED).", "Repo [VERIFY]")
inp("kdp_base", "KDP print cost: fixed part", 1.00, CUR2, "$ per copy", "products/*/listing.json price_notes: $1.00 + $0.07/page premium color.", "Repo [VERIFY]")
inp("kdp_color", "KDP print cost: premium color, per page", 0.07, CUR2, "$ per page", "Same source.", "Repo [VERIFY]")
inp("kdp_bw", "KDP print cost: black-and-white, per page", 0.012, '$0.000', "$ per page", "Not in the repo. Outside figure from general knowledge.", "[VERIFY]")
inp("plays_pages", "100 Screen-Free Plays paperback page count", 82, "0", "pages", "Count of page blocks in products/guide-100-plays/source-kdp.html; confirm at upload.", "Repo [VERIFY]")
inp("pic_pages", "Picture-book and talk-along paperback page count", 32, "0", "pages", "products/*/listing.json ('pages': 32).", "Repo")
inp("ing_disc", "IngramSpark wholesale discount given to retailers", 0.55, PCT1, "% of list", "MARKETING-PLAYBOOK libraries row ('Ingram at 55%'); storefront guide range 30-55%.", "Repo", True)
inp("ing_hc_print", "IngramSpark print cost, 32-page 8.5x8.5 color hardcover", 8.50, CUR2, "$ per copy", "Not in the repo; listing.json says run IngramSpark's calculator.", "[VERIFY]")
inp("ing_pb_print", "IngramSpark print cost, 32-page color paperback", 3.24, CUR2, "$ per copy", "Proxy: KDP premium-color formula from listing.json.", "Assumption [VERIFY]")
inp("mor_pct", "Merchant of record (Gumroad) fee", 0.10, PCT1, "% of price", "storefront-setup-guide §13; legal/international-plan.md (UNVERIFIED). Covers card processing, US sales tax and EU/UK VAT.", "Repo [VERIFY]")
inp("mor_fix", "Merchant of record fixed fee", 0.50, CUR2, "$ per order", "Same source.", "Repo [VERIFY]")
inp("tpt_payout", "Teachers Pay Teachers payout (Basic seller)", 0.55, PCT1, "% of price", "DEMAND-CHECK rule 10; money-and-tax fees table (UNVERIFIED).", "Repo [VERIFY]")
inp("tpt_fix", "TPT per-resource transaction fee", 0.30, CUR2, "$ per sale", "Same source.", "Repo [VERIFY]")
inp("amz_ref", "Amazon Seller Central / FBA referral fee", 0.15, PCT1, "% of price", "storefront-setup-guide §19 ('about 15%', UNVERIFIED).", "Repo [VERIFY]")
inp("fba_fee", "Amazon FBA fulfilment fee per board book", 3.50, CUR2, "$ per unit", "Not in the repo.", "[VERIFY]")
inp("faire_comm", "Faire commission (orders from retailers Faire finds)", 0.15, PCT1, "% of wholesale", "storefront-setup-guide §19; 0% on Faire Direct.", "Repo [VERIFY]")
inp("faire_proc", "Faire payment processing", 0.03, PCT1, "% of wholesale", "Same source; conflicting 1.9-3.5% vs flat ~3%.", "Repo [VERIFY]")
inp("whsl_pct", "Wholesale price as % of retail", 0.50, PCT1, "% of retail", "marketing/AMAZON-AND-RETAIL-ROADMAP.md B4; board-up-go-more/listing.json.", "Repo")
inp("bigbox_ref", "Walmart Marketplace / Target Plus referral fee", 0.15, PCT1, "% of price", "Walmart books about 15% (storefront guide §19, UNVERIFIED). Target Plus is invitation-only; its terms are not in the repo.", "[VERIFY]")
inp("retail_deduct", "Retail deductions and chargebacks allowance", 0.05, PCT1, "% of price", "Not in the repo; big-box programs deduct for compliance misses.", "[VERIFY]")
inp("tee_pod", "Adult tee: POD base cost incl. print", 12.50, CUR2, "$ per unit", "Not in the repo; shipping assumed charged to the buyer.", "[VERIFY]")

a_sec("3. Physical products: board book (Wave 3) and retail stock")
inp("bb_price", "Board book retail price (Up! Go! More!)", 12.99, CUR2, "$", "DEMAND-CHECK §4 rule 9; board-up-go-more/listing.json.", "Repo")
inp("bb_low", "Offset print cost per copy: low", 1.80, CUR2, "$ per copy", "marketing/BLIND-SPOTS.md #11: roughly $1.80-$4 a copy at 1,000-3,000 copies [VERIFY with quotes].", "Repo [VERIFY]")
inp("bb_high", "Offset print cost per copy: high", 4.00, CUR2, "$ per copy", "Same source.", "Repo [VERIFY]")
r = R[0]
put(wa, f"A{r}", "bb_print", font=F_NOTE); put(wa, f"B{r}", "Offset print cost per copy used in model")
put(wa, f"C{r}", f"=C{r-2}+{A['pos']}*(C{r-1}-C{r-2})", fmt=CUR2, font=F_CALC)
put(wa, f"F{r}", "$ per copy"); put(wa, f"G{r}", "Low + cost position x (High - Low)", wrap=True); put(wa, f"H{r}", "Formula")
A["bb_print"] = f"{q('Assumptions')}!$C${r}"; R[0] += 1
inp("freight", "Freight, duties and warehouse receiving", 0.25, PCT1, "% of print cost", "BLIND-SPOTS #11 asks quotes to cover shipping, duties and warehouse fees; no figure given.", "Assumption [VERIFY]")
r = R[0]
put(wa, f"A{r}", "landed", font=F_NOTE); put(wa, f"B{r}", "Landed cost per copy (print + freight/duties)")
put(wa, f"C{r}", f"={A['bb_print']}*(1+{A['freight']})", fmt=CUR2, font=F_CALC)
put(wa, f"F{r}", "$ per copy"); put(wa, f"G{r}", "DEMAND-CHECK rule 9 target: landed cost at or below 35-40% of retail.", wrap=True); put(wa, f"H{r}", "Formula")
A["landed"] = f"{q('Assumptions')}!$C${r}"; R[0] += 1
inp("pick", "3PL pick, pack and ship-handling per single order", 3.50, CUR2, "$ per order", "Not in the repo; postage itself assumed charged to the buyer.", "[VERIFY]")
inp("whs_handle", "3PL handling per unit on wholesale cartons", 1.00, CUR2, "$ per unit", "Not in the repo.", "[VERIFY]")
inp("tpl_min", "3PL monthly minimum / storage", 100, CUR0, "$ per month", "Not in the repo.", "[VERIFY]")
inp("seller_central", "Amazon Seller Central Professional plan", 39.99, CUR2, "$ per month", "storefront-setup-guide §19 (UNVERIFIED).", "Repo [VERIFY]")
inp("min_run", "Minimum first offset run", 1000, NUM0, "copies", "BLIND-SPOTS #11: quotes at 500, 1,000 and 2,500 copies.", "Assumption")
inp("buffer", "Print buffer over pre-sale units", 0.40, PCT1, "%", "BLIND-SPOTS #16: 'Order what sold plus 30-50%'.", "Repo range")
inp("presale_len", "Pre-sale length before the print order", 3, "0", "months", "BLIND-SPOTS #16: Feb-Apr 2027 pre-sale.", "Repo")
inp("presale_buf", "Buffer in the pre-sale funding goal", 0.15, PCT1, "%", "BLIND-SPOTS #16: goal = printing + shipping + duties + fees + delivery + about 15% buffer.", "Repo")
inp("school_setup", "School/group wave one-time setup (TPT fee, purchasing kit, host kit build)", 179, CUR0, "$", "TPT Basic $29 (storefront guide §11); purchasing kit $0-50 and Family Night kit $0-100 (MARKETING-PLAYBOOK).", "Repo")
inp("retail_lead", "Retail readiness spend lead time before retail launch", 3, "0", "months", "Readiness checklist must be done before pitching (AMAZON-AND-RETAIL-ROADMAP B5).", "Assumption")

a_sec("4. Payout delay by channel (months after the sale month)")
CH = [
    ("SITE", "Own site: printables, PDFs and POD tee (Shopify)", "Site sessions (all visitors)", 0, "c",
     "Shopify Payments pays out within days (storefront guide A6)."),
    ("ETSY", "Etsy: printables", "Etsy listing visits", 0, "c", "Etsy deposits on the schedule you choose (A6)."),
    ("KDP", "Amazon KDP: paperbacks", "Amazon detail-page views", 2, "c", "KDP pays about 60 days after month end (A6)."),
    ("INGRAM", "IngramSpark: hardcovers, bookstores, libraries", "Retailer and library listing views (est.)", 3, "c",
     "IngramSpark pays about 90 days after month end (A6, UNVERIFIED)."),
    ("MOR", "International digital via merchant of record (Gumroad)", "International site sessions", 0, "c",
     "Gumroad pays weekly (A6, UNVERIFIED)."),
    ("COURSE", "30-Day Screen Reset written course (Wave 2)", "Subscribers reached by the offer", 0, "c",
     "Sold on Shopify; same payout timing."),
    ("BOARD", "Board book: pre-sale, then 3PL and Amazon FBA (Wave 3)", "Product-page sessions (site + Amazon)", 0, "c",
     "Shopify within days; Seller Central about every 14 days (A6)."),
    ("SCHOOL", "School and group licenses + TPT (held for counsel)", "Schools/orgs page + TPT views", 1, "s",
     "TPT pays monthly; school invoices are net-30 (commerce/PAYMENTS.md)."),
    ("RETAIL", "Retail and wholesale: Faire, Walmart Marketplace, Target Plus", "Retailer and marketplace views", 1, "c",
     "Faire/Walmart settlement timing UNVERIFIED; one month assumed."),
]
CH_NAME = {c[0]: c[1] for c in CH}
for code, name, tl, lag, seas, src in CH:
    inp(f"lag_{code}", f"Payout delay: {name}", lag, "0", "months", src, "Repo" if code not in ("RETAIL",) else "Assumption")

a_sec("5. Seasonality index by calendar month (1.00 = average month)")
a_sub(["ID", "Calendar month", "Consumer index", "School / group index", "", "Unit", "Source / basis", "Status"])
cons = [1.25, 0.90, 1.00, 1.05, 0.95, 1.00, 0.95, 0.85, 0.85, 0.95, 1.30, 1.35]
schl = [1.10, 1.00, 1.00, 0.90, 0.70, 0.40, 0.60, 1.30, 1.40, 1.20, 1.00, 0.70]
mnames = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
s0 = R[0]
for i in range(12):
    r = R[0]
    put(wa, f"A{r}", f"season_{i+1}", font=F_NOTE)
    put(wa, f"B{r}", mnames[i])
    put(wa, f"C{r}", cons[i], fmt=NUM2, font=F_IN)
    put(wa, f"D{r}", schl[i], fmt=NUM2, font=F_IN)
    put(wa, f"F{r}", "index")
    if i == 0:
        put(wa, f"G{r}", "Shape follows the repo calendar: January Reset launch (Wave 2), spring Screen-Free Week, summer kits listed in March, "
                        "holiday gifts Nov-Dec (BLIND-SPOTS #2, #6; DEMAND-CHECK rule 7). School index follows the school year. Levels are assumptions.", wrap=True)
    put(wa, f"H{r}", "Assumption")
    R[0] += 1
A["season_c"] = f"{q('Assumptions')}!$C${s0}:$C${s0+11}"
A["season_s"] = f"{q('Assumptions')}!$D${s0}:$D${s0+11}"

a_sec("6. Scenario drivers: traffic -> conversion -> orders -> order value (yellow = key levers)")
a_sub(["ID", "Driver", "Conservative", "Expected", "Strong", "Unit", "Source / basis", "Status"])
SC = {
    "SITE":   dict(inc=[1, 1, 1], launch=[1, 1, 1], base=[300, 600, 1000], g1=[.08, .12, .12], g2=[.04, .06, .07], g3=[.02, .03, .04], cvr=[.010, .015, .020], items=[1.15, 1.25, 1.35]),
    "ETSY":   dict(inc=[1, 1, 1], launch=[1, 1, 1], base=[400, 800, 1200], g1=[.08, .12, .12], g2=[.03, .05, .06], g3=[.01, .02, .03], cvr=[.010, .020, .025], items=[1.10, 1.20, 1.30]),
    "KDP":    dict(inc=[1, 1, 1], launch=[2, 2, 2], base=[300, 600, 900], g1=[.06, .10, .10], g2=[.03, .04, .05], g3=[.01, .02, .03], cvr=[.03, .05, .06], items=[1.00, 1.05, 1.10]),
    "INGRAM": dict(inc=[1, 1, 1], launch=[3, 3, 3], base=[100, 200, 300], g1=[.04, .06, .06], g2=[.02, .03, .04], g3=[.01, .02, .02], cvr=[.02, .03, .035], items=[1.00, 1.00, 1.05]),
    "MOR":    dict(inc=[1, 1, 1], launch=[4, 4, 4], cvr=[.008, .012, .016], items=[1.10, 1.20, 1.30]),
    "COURSE": dict(inc=[1, 1, 1], launch=[4, 4, 4], cvr=[.010, .015, .020], items=[1, 1, 1]),
    "BOARD":  dict(inc=[1, 1, 1], launch=[5, 5, 5], base=[150, 300, 500], g1=[.04, .06, .06], g2=[.02, .04, .05], g3=[.01, .02, .03], cvr=[.015, .025, .030], items=[1.10, 1.20, 1.30]),
    "SCHOOL": dict(inc=[0, 1, 1], launch=[13, 13, 10], base=[150, 300, 400], g1=[.04, .06, .06], g2=[.04, .06, .06], g3=[.02, .03, .03], cvr=[.010, .015, .020], items=[1.00, 1.10, 1.20]),
    "RETAIL": dict(inc=[0, 1, 1], launch=[25, 25, 22], base=[300, 600, 1200], g1=[.03, .05, .08], g2=[.03, .05, .08], g3=[.03, .05, .08], cvr=[.005, .010, .015], items=[3, 4, 5]),
}
LAUNCH_SRC = {
    "SITE": "Launch-first five live before Black Friday, Nov 27, 2026 (ops/QUEUE.md).",
    "ETSY": "Launch-first five on Etsy + site (DEMAND-CHECK §3).",
    "KDP": "100 Screen-Free Plays paperback, Wave 1 (BLIND-SPOTS #2).",
    "INGRAM": "IngramSpark hardcovers after KDP (storefront guide Part B, step 5).",
    "MOR": "Region 2 English-speaking markets after US launch (ops/INTERNATIONAL.md).",
    "COURSE": "Public New Year launch, January 2027 = month 4 (BLIND-SPOTS #2 Wave 2).",
    "BOARD": "Pre-sale Feb-Apr 2027 = month 5 (BLIND-SPOTS #16; Wave 3).",
    "SCHOOL": "Held until employment counsel clears school-facing work. Conservative assumes it never opens inside 36 months.",
    "RETAIL": "After a 12-month sales record (AMAZON-AND-RETAIL-ROADMAP B1, B5). Conservative assumes no retail inside 36 months.",
}
DRV = [("inc", "Include in scenario (1 = yes, 0 = no)", "0", "flag", True),
       ("launch", "Launch month (model month #)", "0", "month #", False),
       ("base", "Traffic in launch month", NUM0, "visits", True),
       ("g1", "Monthly traffic growth, year 1", PCT1, "% per month", False),
       ("g2", "Monthly traffic growth, year 2", PCT1, "% per month", False),
       ("g3", "Monthly traffic growth, year 3", PCT1, "% per month", False),
       ("cvr", "Conversion rate (visit -> order) at full ramp", PCT2, "% of visits", True),
       ("items", "Items per order", NUM2, "items", False)]
for code, name, tl, lag, seas, src in CH:
    R[0] += 0
    r = R[0]
    for cc in range(1, 9):
        wa.cell(row=r, column=cc).fill = TOT_FILL
    put(wa, f"B{r}", f"{name}  |  traffic = {tl}", bold=True)
    R[0] += 1
    d = SC[code]
    for k, lab, fmt, unit, kl in DRV:
        if k not in d:
            continue
        lab2 = lab
        if code == "RETAIL" and k == "items":
            lab2 = "Units per order (wholesale cartons mixed with single marketplace orders)"
        if code == "COURSE" and k == "cvr":
            lab2 = "Conversion rate (subscriber reached -> purchase)"
        srcx = ""
        status = "Assumption"
        if k == "launch":
            srcx = LAUNCH_SRC[code]
            status = "Repo timing" if code not in ("SCHOOL", "RETAIL") else "Assumption (conditional)"
        elif k == "cvr" and code == "SITE":
            srcx = "MARKETING-PLAYBOOK traffic checklist: working target 1.5-3% site conversion."
            status = "Repo range"
        elif k == "inc" and code in ("SCHOOL", "RETAIL"):
            srcx = "Conditional wave. Set to 0 to see the business without it."
        elif k == "base":
            srcx = "No sales history yet. Replace with the first 90 days of real data."
        inp(f"{code}_{k}", lab2, d[k], fmt, unit, srcx, status, key_lever=kl)
    if code == "SITE":
        inp("intl", "Share of site sessions from outside the US (served by the merchant of record once live)", [.15, .20, .25], PCT1, "% of sessions",
            "International plan: English-speaking Region 2 first (ops/INTERNATIONAL.md).", "Assumption")
    if code == "COURSE":
        inp("reach", "Share of the email list that sees and considers the offer each month", [.10, .15, .20], PCT1, "% of list",
            "Course sold by automated email (DEMAND-CHECK course row).", "Assumption")

R[0] += 1
sec(wa, R[0], 1, 8, "Email list (owned audience; drives the course)")
R[0] += 1
inp("signup", "Site visitors who join the list (free '3 plays for your child's age')", [.02, .03, .04], PCT1, "% of sessions", "BLIND-SPOTS #5 lead magnet.", "Assumption", True)
inp("optin", "Buyers who join through the bonus QR / short link", [.10, .15, .20], PCT1, "% of orders", "BLIND-SPOTS #5: QR code and short link in every product file.", "Assumption")
inp("churn", "Monthly unsubscribe rate", [.020, .015, .010], PCT1, "% of list", "", "Assumption")
R[0] += 1
sec(wa, R[0], 1, 8, "Paid advertising (one test at a time; cut ad sets that miss target CAC)")
R[0] += 1
inp("ads_start", "Ads start month", [2, 2, 2], "0", "month #", "MARKETING-PLAYBOOK week 4 (Oct 19-25): Amazon Ads $5-10/day on live titles.", "Repo")
inp("ads1", "Ad spend per month, year 1", [0, 150, 300], CUR0, "$ per month", "Amazon Ads $5-10/day; Pinterest $10-20/day test after about 60 days (MARKETING-PLAYBOOK).", "Repo range")
inp("ads2", "Ad spend per month, year 2", [150, 300, 600], CUR0, "$ per month", "Adds Google Search at $10/day once site conversion is at least 1.5% (MARKETING-PLAYBOOK #12).", "Repo range")
inp("ads3", "Ad spend per month, year 3", [250, 450, 900], CUR0, "$ per month", "Raise only while CAC is under target (MARKETING-PLAYBOOK checklist).", "Assumption")

wa.column_dimensions["A"].width = 14
wa.column_dimensions["B"].width = 58
for c in "CDE":
    wa.column_dimensions[c].width = 14
wa.column_dimensions["F"].width = 14
wa.column_dimensions["G"].width = 70
wa.column_dimensions["H"].width = 16
wa.freeze_panes = "C6"
dv = DataValidation(type="whole", operator="between", formula1="0", formula2="1", allow_blank=False)
wa.add_data_validation(dv)
for code in ("SITE", "ETSY", "KDP", "INGRAM", "MOR", "COURSE", "BOARD", "SCHOOL", "RETAIL"):
    for ref in A[f"{code}_inc"]:
        dv.add(ref.split("!")[1].replace("$", ""))
dvp = DataValidation(type="decimal", operator="between", formula1="0", formula2="1")
wa.add_data_validation(dvp)
dvp.add(A["pos"].split("!")[1].replace("$", ""))

# =====================================================================
# UNIT ECONOMICS
# =====================================================================
wu = ws_["Unit Economics"]
title(wu, "Unit economics: net to AlphaPlay LLC per unit, by product and channel",
      "Net per unit = Price x (1 - platform % - payment % - refund %) - fixed fees - print/POD/landed cost - fulfilment. "
      "Sales tax and VAT are pass-through and never counted as revenue. Blue prices are editable; other channels link to them.")
wu.merge_cells("A2:R3"); wu["A2"].alignment = WRAP; wu.row_dimensions[2].height = 28
UH = ["Channel code", "Product", "Channel / route", "Wave", "Price to customer ($)", "Platform / marketplace fee %",
      "Platform fixed fee ($/unit)", "Payment processing %", "Payment fixed fee ($)", "Refund / deduction allowance %",
      "Print, POD or landed cost ($/unit)", "Fulfilment ($/unit)", "Net per unit ($)", "Net margin %",
      "Inventory cost paid up front ($/unit)", "Mix weight in channel", "Sales tax / VAT handling", "Source / notes"]
hdr(wu, 5, 1, 18, UH, 45)
a = A
ETSY_F = f"={a['etsy_tx']}+{a['etsy_oa']}*{a['etsy_oashare']}"
TAX_SHOP = "Shopify Tax adds MD 6% at checkout; AlphaPlay remits (pass-through, not revenue)"
TAX_ETSY = "Etsy is marketplace facilitator: collects and remits"
TAX_KDP = "Amazon is the retailer; royalty only, no sales tax for AlphaPlay"
TAX_ING = "Retailer collects; AlphaPlay receives publisher compensation"
TAX_MOR = "Merchant of record collects and remits US sales tax and EU/UK VAT"
TAX_TPT = "TPT remits as marketplace (UNVERIFIED)"
TAX_SCH = "Tax-exempt certificate kept on file; otherwise MD 6%"
TAX_AMZ = "Amazon collects and remits"
TAX_FAIRE = "Wholesale: exempt with the retailer's resale certificate"
TAX_BB = "Marketplace facilitator [VERIFY for Target Plus]"
SHOP = dict(F=0, G=0, H=a["shop_pct"], I=a["shop_fix"], J=a["refund"], tax=TAX_SHOP)
ETSY = dict(F=ETSY_F, G=a["etsy_list"], H=a["etsy_pay"], I=a["etsy_payfix"], J=a["refund"], tax=TAX_ETSY)
MOR = dict(F=a["mor_pct"], G=a["mor_fix"], H=0, I=0, J=a["refund"], tax=TAX_MOR)
KDPc = dict(F=f"=1-{a['kdp_roy']}", G=0, H=0, I=0, J=0, tax=TAX_KDP)
ING = dict(F=a["ing_disc"], G=0, H=0, I=0, J=0, tax=TAX_ING)
TPT = dict(F=f"=1-{a['tpt_payout']}", G=a["tpt_fix"], H=0, I=0, J=0, tax=TAX_TPT)
SCHD = dict(F=0, G=0, H=a["shop_pct"], I=a["shop_fix"], J=0, tax=TAX_SCH)

KDP_COLOR = f"={a['kdp_base']}+{a['kdp_color']}*{a['pic_pages']}"
KDP_BW = f"={a['kdp_base']}+{a['kdp_bw']}*{a['plays_pages']}"
DIGI = [  # id, product, price, weight, src
    ("routine", "Visual routine cards (200+ editable, 0-5 and 5-12)", 6.50, 25, "DEMAND-CHECK §3 #1: $9.50 list / about $6.50 sale. Everyday price set at $6.50 under BRAND.md 'Honest pricing' (no permanent anchor discounts)."),
    ("bored", "\"I'm bored\" play cards (150, age-banded)", 6.50, 15, "DEMAND-CHECK §3 #2."),
    ("family", "Play-First Family Kit", 11.00, 20, "DEMAND-CHECK §3 #3."),
    ("busy", "Toddler busy book printable (120-150 pages)", 11.99, 18, "DEMAND-CHECK §3 #4: $15.99 list, sold at about $11-12; everyday price $11.99 under Honest pricing."),
    ("playspdf", "100 Screen-Free Plays (PDF)", 9.99, 10, "DEMAND-CHECK §3 #5."),
    ("car", "Screen-Free Car Ride & Waiting Pack", 6.00, 5, "DEMAND-CHECK §3 first alternate."),
    ("flash", "First-words flash cards (printable)", 6.99, 3, "DEMAND-CHECK §2 #3."),
    ("alpha", "ALPHAPLAY Spelling Games printable", 6.00, 2, "ops/QUEUE.md #1 (trademark use). No price in repo; $6 is an assumption (never list a single under $5, rule 3)."),
]
UROWS = []  # dicts
for pid, pname, price, w, src in DIGI:
    UROWS.append(dict(code="SITE", pid=f"SITE_{pid}", prod=pname, route="Own site (Shopify)", wave="1", price=price, w=w, K=0, L=0, O=0, src=src, **SHOP))
UROWS.append(dict(code="SITE", pid="SITE_tee", prod="Adult tee (POD, 2-3 cleared designs max)", route="Own site (Shopify) + POD partner", wave="1",
                  price=27, w=2, K=a["tee_pod"], L=0, O=0, src="DEMAND-CHECK tees row: $27; no merch line beyond 2-3 designs.", **SHOP))
for pid, pname, price, w, src in DIGI:
    UROWS.append(dict(code="ETSY", pid=f"ETSY_{pid}", prod=pname, route="Etsy", wave="1", price=("ref", f"SITE_{pid}"), w=w, K=0, L=0, O=0,
                      src="Price linked to the own-site row.", **ETSY))
for pid, pname, price, w, src in DIGI[:5]:
    UROWS.append(dict(code="MOR", pid=f"MOR_{pid}", prod=pname, route="Gumroad (merchant of record), buyers outside the US", wave="1-2",
                      price=("ref", f"SITE_{pid}"), w=w, K=0, L=0, O=0, src="International plan: one MoR for international digital sales.", **MOR))
UROWS += [
    dict(code="KDP", pid="KDP_plays", prod="100 Screen-Free Plays (paperback, B/W interior)", route="Amazon KDP", wave="1", price=16.99, w=50,
         K=KDP_BW, L=0, O=0, src="DEMAND-CHECK §3 #5: $16.99 paperback. Print cost uses the B/W rate [VERIFY].", **KDPc),
    dict(code="KDP", pid="KDP_tablet", prod="The Day the Tablet Slept (paperback)", route="Amazon KDP", wave="1", price=11.99, w=20,
         K=KDP_COLOR, L=0, O=0, src="picture-tablet-slept/listing.json: $11.99; about $3.95/copy royalty.", **KDPc),
    dict(code="KDP", pid="KDP_upgo", prod="Up! Go! More! (talk-along paperback)", route="Amazon KDP", wave="1", price=11.99, w=30,
         K=KDP_COLOR, L=0, O=0, src="board-up-go-more/listing.json: $11.99 paperback edition sells now with no inventory.", **KDPc),
    dict(code="INGRAM", pid="ING_laps", prod="Laps Not Apps (32-page hardcover)", route="IngramSpark", wave="1", price=19.99, w=35,
         K=a["ing_hc_print"], L=0, O=0, src="picture-laps-not-apps/listing.json: $19.99; aim print cost at or below 35-40% of list.", **ING),
    dict(code="INGRAM", pid="ING_tablet", prod="The Day the Tablet Slept (hardcover)", route="IngramSpark", wave="1", price=19.99, w=25,
         K=a["ing_hc_print"], L=0, O=0, src="picture-tablet-slept/listing.json: hardcover $19.99.", **ING),
    dict(code="INGRAM", pid="ING_plays", prod="100 Screen-Free Plays (paperback via Ingram)", route="IngramSpark (same ISBN; KDP Expanded Distribution off)", wave="1",
         price=("ref", "KDP_plays"), w=20, K=KDP_BW, L=0, O=0, src="storefront-setup-guide Part D. Print cost proxy = KDP B/W formula [VERIFY].", **ING),
    dict(code="INGRAM", pid="ING_upgo", prod="Up! Go! More! (paperback via Ingram)", route="IngramSpark", wave="1",
         price=("ref", "KDP_upgo"), w=20, K=a["ing_pb_print"], L=0, O=0, src="board-up-go-more/listing.json channels.", **ING),
    dict(code="COURSE", pid="CRS_single", prod="30-Day Screen Reset (written program)", route="Own site (Shopify) + automated email", wave="2", price=27, w=70,
         K=0, L=0, O=0, src="DEMAND-CHECK course row: $27.", **dict(SHOP, J=a["refund_course"])),
    dict(code="COURSE", pid="CRS_bundle", prod="30-Day Screen Reset bundle", route="Own site (Shopify) + automated email", wave="2", price=49, w=30,
         K=0, L=0, O=0, src="DEMAND-CHECK course row: $49 bundle.", **dict(SHOP, J=a["refund_course"])),
    dict(code="BOARD", pid="BB_site", prod="Up! Go! More! board book", route="Own site; offset stock at a 3PL", wave="3", price=("a", a["bb_price"]), w=60,
         K=a["landed"], L=a["pick"], O=a["landed"], src="Offset run held and shipped by a 3PL, never the founder (BRAND.md). Postage charged to the buyer.", **SHOP),
    dict(code="BOARD", pid="BB_fba", prod="Up! Go! More! board book", route="Amazon FBA", wave="3", price=("a", a["bb_price"]), w=40,
         K=a["landed"], L=a["fba_fee"], O=a["landed"], src="AMAZON-AND-RETAIL-ROADMAP A: board books via FBA; CPSIA paperwork first.",
         F=a["amz_ref"], G=0, H=0, I=0, J=0, tax=TAX_AMZ),
    dict(code="SCHOOL", pid="SCH_tpt", prod="TPT single resource (e.g., talk brain breaks)", route="Teachers Pay Teachers", wave="4 (held)", price=5, w=40,
         K=0, L=0, O=0, src="DEMAND-CHECK teacher table: $5 singles; held for counsel.", **TPT),
    dict(code="SCHOOL", pid="SCH_tptb", prod="Classroom Talk and Play growth bundle", route="Teachers Pay Teachers", wave="4 (held)", price=22, w=20,
         K=0, L=0, O=0, src="DEMAND-CHECK §2 #11: $22-26.", **TPT),
    dict(code="SCHOOL", pid="SCH_class", prod="Single-classroom license", route="Direct invoice / Shopify draft order", wave="4 (held)", price=29, w=20,
         K=0, L=0, O=0, src="MARKETING-PLAYBOOK School Purchasing Kit: Single Classroom $29.", **SCHD),
    dict(code="SCHOOL", pid="SCH_kit", prod="Host-it-yourself parent-night kit, single site", route="Direct invoice / Shopify", wave="4 (held)", price=129, w=15,
         K=0, L=0, O=0, src="DEMAND-CHECK: $129 single site.", **SCHD),
    dict(code="SCHOOL", pid="SCH_kitm", prod="Host kit, multi-site", route="Direct invoice / Shopify", wave="4 (held)", price=249, w=5,
         K=0, L=0, O=0, src="DEMAND-CHECK: $249 multi-site.", **SCHD),
    dict(code="RETAIL", pid="RT_faire", prod="Up! Go! More! board book (wholesale)", route="Faire wholesale to boutiques", wave="5 (retail)",
         price=("f", f"={a['bb_price']}*{a['whsl_pct']}"), w=50, K=a["landed"], L=a["whs_handle"], O=a["landed"],
         src="Wholesale pays about half of retail (AMAZON-AND-RETAIL-ROADMAP B4).", F=a["faire_comm"], G=0, H=a["faire_proc"], I=0, J=a["retail_deduct"], tax=TAX_FAIRE),
    dict(code="RETAIL", pid="RT_wmt", prod="Up! Go! More! board book", route="Walmart Marketplace (3PL-fulfilled)", wave="5 (retail)",
         price=("a", a["bb_price"]), w=25, K=a["landed"], L=a["pick"], O=a["landed"], src="Walmart Marketplace by application (ROADMAP B3).",
         F=a["bigbox_ref"], G=0, H=0, I=0, J=a["retail_deduct"], tax=TAX_BB),
    dict(code="RETAIL", pid="RT_tgt", prod="Up! Go! More! board book", route="Target Plus (invitation-only) [VERIFY]", wave="5 (retail)",
         price=("a", a["bb_price"]), w=25, K=a["landed"], L=a["pick"], O=a["landed"], src="Target Plus is invitation-only (ROADMAP B3) [VERIFY terms].",
         F=a["bigbox_ref"], G=0, H=0, I=0, J=a["retail_deduct"], tax=TAX_BB),
]
UE_ROW = {}
r = 6
for d in UROWS:
    UE_ROW[d["pid"]] = r
    r += 1
UE_FIRST, UE_LAST = 6, r - 1

def fval(v):
    if isinstance(v, str) and v.startswith("'"):
        return "=" + v
    return v

for d in UROWS:
    r = UE_ROW[d["pid"]]
    put(wu, f"A{r}", d["code"])
    put(wu, f"B{r}", d["prod"])
    put(wu, f"C{r}", d["route"])
    put(wu, f"D{r}", d["wave"])
    p = d["price"]
    if isinstance(p, tuple):
        if p[0] == "ref":
            put(wu, f"E{r}", f"=E{UE_ROW[p[1]]}", fmt=CUR2)
        elif p[0] == "a":
            put(wu, f"E{r}", f"={p[1]}", fmt=CUR2)
        else:
            put(wu, f"E{r}", p[1], fmt=CUR2)
    else:
        put(wu, f"E{r}", p, fmt=CUR2, font=F_IN)
    for colk, fmt in (("F", PCT1), ("G", CUR2), ("H", PCT1), ("I", CUR2), ("J", PCT1), ("K", CUR2), ("L", CUR2), ("O", CUR2)):
        v = fval(d[colk])
        if isinstance(v, (int, float)):
            put(wu, f"{colk}{r}", v, fmt=fmt, font=F_IN if v != 0 else F_CALC)
        else:
            put(wu, f"{colk}{r}", v, fmt=fmt)
    put(wu, f"M{r}", f"=E{r}*(1-F{r}-H{r}-J{r})-G{r}-I{r}-K{r}-L{r}", fmt=CUR2)
    put(wu, f"N{r}", f"=IF(E{r}=0,0,M{r}/E{r})", fmt=PCT1)
    put(wu, f"P{r}", d["w"], fmt="0", font=F_IN)
    put(wu, f"Q{r}", d["tax"], wrap=False)
    put(wu, f"R{r}", d["src"], wrap=False)

# channel summary
rs = UE_LAST + 3
sec(wu, rs - 1, 1, 18, "Channel blend used by the forecast (weighted by mix weight)")
hdr(wu, rs, 1, 9, ["Channel code", "Channel", "", "", "Blended price ($/item)", "Blended net per item ($)", "Net margin %",
                   "Inventory cost paid up front ($/item)", "Payout delay (months)"], 32)
UE = {}
rng = lambda col: f"${col}${UE_FIRST}:${col}${UE_LAST}"
for i, (code, name, tl, lag, seas, src) in enumerate(CH):
    r = rs + 1 + i
    put(wu, f"A{r}", code, bold=True)
    put(wu, f"B{r}", name)
    put(wu, f"E{r}", f"=SUMPRODUCT(--({rng('A')}=$A{r}),{rng('E')},{rng('P')})/SUMIF({rng('A')},$A{r},{rng('P')})", fmt=CUR2)
    put(wu, f"F{r}", f"=SUMPRODUCT(--({rng('A')}=$A{r}),{rng('M')},{rng('P')})/SUMIF({rng('A')},$A{r},{rng('P')})", fmt=CUR2)
    put(wu, f"G{r}", f"=IF(E{r}=0,0,F{r}/E{r})", fmt=PCT1)
    put(wu, f"H{r}", f"=SUMPRODUCT(--({rng('A')}=$A{r}),{rng('O')},{rng('P')})/SUMIF({rng('A')},$A{r},{rng('P')})", fmt=CUR2)
    put(wu, f"I{r}", f"={A['lag_'+code]}", fmt="0")
    UE[code] = dict(price=f"{q('Unit Economics')}!$E${r}", net=f"{q('Unit Economics')}!$F${r}",
                    landed=f"{q('Unit Economics')}!$H${r}", row=r)
rn = rs + len(CH) + 2
notes = [
    "Notes",
    "KDP and IngramSpark: 'price' is the list price the reader pays; AlphaPlay receives the royalty or publisher compensation shown in Net per unit.",
    "Board book and retail rows: the landed print cost is deducted in Net per unit (true margin). Cash Flow adds it back and instead pays for whole print runs up front.",
    "Fixed per-order fees are applied per item, which slightly understates net on multi-item orders (a conservative simplification).",
    "IngramSpark hardcover margins are thin at a 55% discount; test a 40% discount or a higher list price (listing.json asks each sale to clear $2-$3).",
    "Postage on physical orders is assumed charged to the buyer at cost (pass-through).",
]
for i, t in enumerate(notes):
    put(wu, f"A{rn+i}", t, font=fnt(bold=(i == 0)) if i == 0 else F_NOTE)
widths = [10, 44, 34, 9, 11, 11, 11, 11, 11, 11, 12, 11, 11, 10, 12, 9, 52, 80]
for i, w in enumerate(widths):
    wu.column_dimensions[CL(i + 1)].width = w
wu.freeze_panes = "C6"

# =====================================================================
# STARTUP COSTS
# =====================================================================
wsu = ws_["Startup Costs"]
title(wsu, "Startup costs: one-time items",
      "Model ($) = Low + cost position x (High - Low), using the cost position on Assumptions. Set Include to 0 to drop an item. "
      "Recurring items (insurance, domains, subscriptions) live on Monthly Operating Costs.")
wsu.merge_cells("A2:I3"); wsu["A2"].alignment = WRAP; wsu.row_dimensions[2].height = 28
hdr(wsu, 5, 1, 9, ["Item", "Category", "Low ($)", "High ($)", "Model ($)", "Month incurred (model month #)", "Include (1/0)", "Source", "Status"], 32)
SU = [
    ("Trade-name filing, 'Play Before Pixels' (SDAT)", "Legal/entity", 25, 50, 1, 1, "legal/protection/PROTECTION-PLAN.md: $25 ($50 expedited).", "Repo"),
    ("Principal office / resident agent change (SDAT)", "Legal/entity", 25, 50, 1, 1, "PROTECTION-PLAN: $25 ($50 expedited).", "Repo"),
    ("DMCA designated agent registration", "Legal/IP", 6, 6, 1, 1, "PROTECTION-PLAN: $6, renews every 3 years.", "Repo"),
    ("Etsy shop opening fee (may apply)", "Channels", 0, 15, 1, 1, "storefront-setup-guide §10.", "Repo [VERIFY]"),
    ("Bowker ISBNs, 10-pack", "Publishing", 295, 295, 1, 1, "PROTECTION-PLAN; DECISION-MEMO (UNVERIFIED).", "Repo [VERIFY]"),
    ("Trademark clearance search + attorney opinion, PLAY BEFORE PIXELS", "Legal/IP", 500, 2500, 1, 1, "legal/DECISION-MEMO.json step 2.", "Repo [VERIFY]"),
    ("USPTO application, PLAY BEFORE PIXELS, classes 16 + 41", "Legal/IP", 700, 700, 2, 1, "PROTECTION-PLAN #7: $350/class, within 90 days.", "Repo"),
    ("ALPHAPLAY Statement of Use ($150/class) or extension (up to $625)", "Legal/IP", 150, 625, 5, 1, "PROTECTION-PLAN #6: due Mar 8, 2027; target Feb 1, 2027.", "Repo"),
    ("Trademark attorney, ALPHAPLAY filing", "Legal/IP", 300, 1000, 4, 1, "marketing/BLIND-SPOTS.md #9.", "Repo"),
    ("Copyright registrations for pending works (before Nov 11, 2026)", "Legal/IP", 360, 680, 2, 1, "PROTECTION-PLAN: about 8 filings at $45-$85.", "Repo"),
    ("Attorney: IP assignment to the LLC + operating-agreement refresh", "Legal/entity", 500, 2000, 3, 1, "PROTECTION-PLAN #5 (after employment counsel answers).", "Repo [VERIFY]"),
    ("Attorney: Terms of Sale and website legal pages", "Legal", 500, 2500, 1, 1, "PROTECTION-PLAN contracts table; DECISION-MEMO $500-$2,000.", "Repo [VERIFY]"),
    ("Attorney: freelancer work-for-hire template", "Legal", 300, 1000, 5, 1, "PROTECTION-PLAN contracts table.", "Repo [VERIFY]"),
    ("Attorney: digital product license review", "Legal", 0, 300, 1, 1, "PROTECTION-PLAN contracts table.", "Repo [VERIFY]"),
    ("Attorney: successor in operating agreement + durable POA", "Legal/entity", 500, 1500, 6, 1, "BLIND-SPOTS #20.", "Repo [VERIFY]"),
    ("Accountant setup review (W-9 answer, sales tax, chart of accounts)", "Finance", 300, 1000, 1, 1, "No fee in the repo (finance/money-and-tax-setup.md lists the questions).", "Assumption [VERIFY]"),
    ("Print proofs, about 6 print products at $5-$15", "Product", 30, 90, 1, 1, "BLIND-SPOTS #11.", "Repo"),
    ("Expert accuracy review (CCC-SLP, flat fee)", "Product", 200, 500, 3, 1, "BLIND-SPOTS #10.", "Repo"),
    ("Sensitivity read", "Product", 150, 600, 3, 1, "BLIND-SPOTS #10.", "Repo"),
    ("Early review copies for review teams", "Marketing", 100, 300, 2, 1, "BLIND-SPOTS #14.", "Repo"),
    ("Holiday gift setup (gift cards, reveal card)", "Marketing", 0, 20, 1, 1, "BLIND-SPOTS #6.", "Repo"),
    ("Human illustrator for the board book (Wave 3)", "Product (Wave 3)", 1500, 5000, 6, 1, "BLIND-SPOTS #17.", "Repo"),
    ("Library cataloging block", "Library credibility", 75, 150, 9, 1, "BLIND-SPOTS #18 [VERIFY].", "Repo [VERIFY]"),
    ("One paid review (Kirkus Indie or Foreword Clarion)", "Library credibility", 500, 650, 9, 1, "BLIND-SPOTS #18 [VERIFY].", "Repo [VERIFY]"),
    ("Juried award entries (2-3 at $75-$100)", "Library credibility", 225, 300, 10, 1, "BLIND-SPOTS #18 [VERIFY].", "Repo [VERIFY]"),
    ("Spanish 'starter' localization (AI + professional post-edit)", "International", 4700, 5800, 7, 1, "legal/international-plan.md: Spanish mixed, months 4-9.", "Repo (estimate)"),
    ("Madrid international trademark (EU, UK, CA, AU; 2 classes) - optional", "International", 2500, 4500, 8, 0, "DECISION-MEMO: only once the US application looks safe.", "Repo [VERIFY]"),
]
for i, (item, cat, lo, hi, m, inc, src, st) in enumerate(SU):
    r = 6 + i
    put(wsu, f"A{r}", item); put(wsu, f"B{r}", cat)
    put(wsu, f"C{r}", lo, fmt=CUR0, font=F_IN); put(wsu, f"D{r}", hi, fmt=CUR0, font=F_IN)
    put(wsu, f"E{r}", f"=C{r}+{A['pos']}*(D{r}-C{r})", fmt=CUR0)
    put(wsu, f"F{r}", m, fmt="0", font=F_IN); put(wsu, f"G{r}", inc, fmt="0", font=F_IN)
    put(wsu, f"H{r}", src); put(wsu, f"I{r}", st)
SU_FIRST, SU_LAST = 6, 6 + len(SU) - 1
rt = SU_LAST + 1
put(wsu, f"A{rt}", "Total (included items)", bold=True)
put(wsu, f"C{rt}", f"=SUMPRODUCT(C{SU_FIRST}:C{SU_LAST},$G${SU_FIRST}:$G${SU_LAST})", fmt=CUR0, bold=True)
put(wsu, f"D{rt}", f"=SUMPRODUCT(D{SU_FIRST}:D{SU_LAST},$G${SU_FIRST}:$G${SU_LAST})", fmt=CUR0, bold=True)
put(wsu, f"E{rt}", f"=SUMPRODUCT(E{SU_FIRST}:E{SU_LAST},$G${SU_FIRST}:$G${SU_LAST})", fmt=CUR0, bold=True)
for c in "ABCDEFGHI":
    wsu[f"{c}{rt}"].border = TOP; wsu[f"{c}{rt}"].fill = TOT_FILL
put(wsu, f"A{rt+2}", "Timing summary", bold=True)
for j, (lab, lo_m, hi_m) in enumerate([("Months 1-3 (Oct-Dec 2026)", 1, 3), ("Months 4-12", 4, 12), ("Year 2", 13, 24), ("Year 3", 25, 36)]):
    r = rt + 3 + j
    put(wsu, f"A{r}", lab)
    put(wsu, f"E{r}", f"=SUMPRODUCT(($F${SU_FIRST}:$F${SU_LAST}>={lo_m})*($F${SU_FIRST}:$F${SU_LAST}<={hi_m})*$G${SU_FIRST}:$G${SU_LAST}*$E${SU_FIRST}:$E${SU_LAST})", fmt=CUR0)
put(wsu, f"A{rt+8}", "Not included: employment-counsel fees (a personal matter for the founder); board-book CPSIA testing, 3PL setup and print runs (see Retail Readiness Costs and Cash Flow).", font=F_NOTE)
for c, w in zip("ABCDEFGHI", [62, 18, 11, 11, 11, 14, 10, 62, 18]):
    wsu.column_dimensions[c].width = w
wsu.freeze_panes = "B6"
SU_RANGE = dict(F=f"{q('Startup Costs')}!$F${SU_FIRST}:$F${SU_LAST}", G=f"{q('Startup Costs')}!$G${SU_FIRST}:$G${SU_LAST}",
                E=f"{q('Startup Costs')}!$E${SU_FIRST}:$E${SU_LAST}")

# =====================================================================
# MONTHLY OPERATING COSTS
# =====================================================================
wo = ws_["Monthly Operating Costs"]
title(wo, "Monthly operating costs: subscriptions, Claude plan, email platform, connectors, insurance and reserve",
      "Low/High are the amount per billing period (per month for Monthly items, per year for Annual items). Model = Low + cost position x (High - Low). "
      "Annual items hit in their start month and every 12 months after. Insurance is shown billed monthly (annual premium / 12).")
wo.merge_cells("A2:M3"); wo["A2"].alignment = WRAP; wo.row_dimensions[2].height = 28
GC = 15  # grid first column (O)
hdr(wo, 5, 1, 13, ["Item", "Category", "Low ($)", "High ($)", "Model ($ per period)", "Frequency", "Start month", "End month",
                   "Year 1 total", "Year 2 total", "Year 3 total", "Source", "Status"], 32)
for m in range(1, 37):
    c = CL(GC + m - 1)
    put(wo, f"{c}3", f"=INT(({c}4-1)/12)+1", fmt="0")
    put(wo, f"{c}4", m, fmt="0")
    put(wo, f"{c}5", f"=DATE(YEAR({A['start']}),MONTH({A['start']})+{c}4-1,1)", fmt=DATEF)
    wo[f"{c}5"].fill = HDR_FILL; wo[f"{c}5"].font = HDR_FONT
    wo.column_dimensions[c].width = 9
put(wo, f"{CL(GC-1)}3", "Year", font=F_NOTE); put(wo, f"{CL(GC-1)}4", "Month #", font=F_NOTE)
OP = [
    ("Shopify Basic plan", "Store", 39, 39, "Monthly", 1, 36, "storefront-setup-guide §1: $39/mo ($29 billed yearly).", "Repo [VERIFY]"),
    ("Business email on the brand domain", "Tools", 7, 14, "Monthly", 1, 36, "BLIND-SPOTS #4: about $7-$14 a month.", "Repo"),
    ("Password manager", "Tools", 0, 40, "Annual", 1, 36, "BLIND-SPOTS #4: $0-$40 a year.", "Repo"),
    ("Email platform, year 1 (free tiers, then budget)", "Email", 0, 20, "Monthly", 1, 12, "MARKETING-PLAYBOOK: Klaviyo free 250 profiles; budget $20/mo.", "Repo [VERIFY]"),
    ("Email platform, years 2-3", "Email", 20, 39, "Monthly", 13, 36, "MARKETING-PLAYBOOK: about $20-39/mo later.", "Repo [VERIFY]"),
    ("Claude plan that runs the scheduled routines", "Automation", 100, 200, "Monthly", 1, 36, "No price in the repo; a higher-usage tier is assumed because routines run weekly (ops/ROUTINE.md).", "Assumption [VERIFY]"),
    ("Claude usage-credit cap (hard limit)", "Automation", 0, 25, "Monthly", 1, 36, "ops/GAPS-ROUND-2.md G2-06: fixed monthly limit such as $25.", "Repo"),
    ("QuickBooks Online (Simple Start to Essentials)", "Bookkeeping", 38, 85, "Monthly", 1, 36, "finance/money-and-tax-setup.md: $38 Simple Start; Essentials $75-$85 (UNVERIFIED).", "Repo [VERIFY]"),
    ("Link My Books settlement connector (1-3 channels)", "Connectors", 21, 47, "Monthly", 3, 36, "money-and-tax-setup: from $21/mo, about $13 per extra channel.", "Repo [VERIFY]"),
    ("Buyer-specific PDF stamping", "Tools", 0, 20, "Monthly", 1, 36, "PROTECTION-PLAN: about $0-$20/mo.", "Repo [VERIFY]"),
    ("Etsy listing renewals (about 30 listings every 4 months)", "Channels", 1.5, 3, "Monthly", 1, 36, "$0.20 per listing (storefront guide §10); listing count is an assumption.", "Assumption"),
    ("Sales-tax filing service (optional; the accountant may file)", "Tax", 0, 25, "Monthly", 1, 36, "TAX-AUTOPILOT §1 names services; no price in the repo.", "Assumption [VERIFY]"),
    ("Domains, MUST tier", "Domains", 61, 61, "Annual", 1, 36, "legal/domain-portfolio.md: about $61 a year.", "Repo [VERIFY]"),
    ("Domains, SHOULD tier (before paid ads / first print run)", "Domains", 118, 118, "Annual", 5, 36, "legal/domain-portfolio.md: about $118 a year.", "Repo [VERIFY]"),
    ("USPS PO Box (public business address)", "Admin", 100, 300, "Annual", 1, 36, "legal/ENTITY.md decision; fee not in the repo (TRUST-CHECKLIST says verify at usps.com).", "Assumption [VERIFY]"),
    ("Commercial resident agent", "Admin", 50, 300, "Annual", 1, 36, "PROTECTION-PLAN: about $50-$300 a year.", "Repo [VERIFY]"),
    ("Maryland SDAT annual report", "Admin", 300, 300, "Annual", 7, 36, "PROTECTION-PLAN: $300, due Apr 15 (month 7 = Apr 2027).", "Repo"),
    ("General liability with products-completed operations", "Insurance", 45.17, 125, "Monthly", 1, 36, "PROTECTION-PLAN: about $542/yr average; range $260-$3,000+ (low = average / 12).", "Repo [VERIFY]"),
    ("Professional liability / E&O", "Insurance", 62, 125, "Monthly", 1, 36, "PROTECTION-PLAN: about $62/mo average; $400-$3,750/yr range.", "Repo [VERIFY]"),
    ("Media liability / publisher's E&O", "Insurance", 50, 150, "Monthly", 1, 36, "PROTECTION-PLAN: $50-$150/mo (secondary, unverified).", "Repo [VERIFY]"),
    ("Cyber", "Insurance", 35, 129, "Monthly", 1, 36, "PROTECTION-PLAN: $35-$129/mo.", "Repo [VERIFY]"),
    ("Umbrella / excess over GL", "Insurance", 25, 50, "Monthly", 3, 36, "PROTECTION-PLAN: 'a few hundred dollars a year' (unverified).", "Repo [VERIFY]"),
    ("Trademark watch (optional; the monthly DIY search is $0)", "Legal/IP", 0, 750, "Annual", 4, 36, "PROTECTION-PLAN: $395-$750 a year.", "Repo [VERIFY]"),
    ("Accountant: year-end return and review", "Finance", 500, 1500, "Annual", 7, 36, "No fee in the repo.", "Assumption [VERIFY]"),
    ("Copyright registrations for new releases (about 4 a year)", "Legal/IP", 260, 340, "Annual", 13, 36, "PROTECTION-PLAN: $65-$85 per filing.", "Repo"),
]
OP_FIRST = 6
for i, (item, cat, lo, hi, fq, st, en, src, stt) in enumerate(OP):
    r = OP_FIRST + i
    put(wo, f"A{r}", item); put(wo, f"B{r}", cat)
    put(wo, f"C{r}", lo, fmt=CUR2, font=F_IN); put(wo, f"D{r}", hi, fmt=CUR2, font=F_IN)
    put(wo, f"E{r}", f"=C{r}+{A['pos']}*(D{r}-C{r})", fmt=CUR2)
    put(wo, f"F{r}", fq, font=F_IN); put(wo, f"G{r}", st, fmt="0", font=F_IN); put(wo, f"H{r}", en, fmt="0", font=F_IN)
    g1, gl = CL(GC), CL(GC + 35)
    for j, col in enumerate("IJK"):
        put(wo, f"{col}{r}", f"=SUMIF(${g1}$3:${gl}$3,{j+1},{g1}{r}:{gl}{r})", fmt=CUR0)
    put(wo, f"L{r}", src); put(wo, f"M{r}", stt)
    for m in range(1, 37):
        c = CL(GC + m - 1)
        put(wo, f"{c}{r}", f"=IF(AND({c}$4>=$G{r},{c}$4<=$H{r}),IF($F{r}=\"Monthly\",$E{r},IF(MOD({c}$4-$G{r},12)=0,$E{r},0)),0)", fmt=CUR0)
OP_LAST = OP_FIRST + len(OP) - 1
r_sub, r_con, r_tot = OP_LAST + 1, OP_LAST + 2, OP_LAST + 3
labels = {r_sub: "Subtotal", r_con: "Operating contingency", r_tot: "Total monthly operating costs"}
for rr in (r_sub, r_con, r_tot):
    put(wo, f"A{rr}", labels[rr], bold=True)
    for c in range(1, GC + 36):
        wo.cell(row=rr, column=c).fill = TOT_FILL
    for colidx in list(range(9, 12)) + list(range(GC, GC + 36)):
        c = CL(colidx)
        if rr == r_sub:
            f = f"=SUM({c}{OP_FIRST}:{c}{OP_LAST})"
        elif rr == r_con:
            f = f"={c}{r_sub}*{A['contin']}"
        else:
            f = f"={c}{r_sub}+{c}{r_con}"
        put(wo, f"{c}{rr}", f, fmt=CUR0, bold=(rr == r_tot))
wo[f"A{r_sub}"].border = TOP
put(wo, f"A{r_tot+2}", "Wave-specific fixed costs (3PL minimum, Seller Central, retail EDI and insurance uplift) are added in the forecast only when that wave is live.", font=F_NOTE)
for c, w in zip("ABCDEFGHIJKLM", [52, 12, 10, 10, 12, 10, 8, 8, 11, 11, 11, 60, 18]):
    wo.column_dimensions[c].width = w
wo.column_dimensions["N"].width = 8
wo.freeze_panes = f"{CL(GC)}6"
OPTOT = lambda m: f"{q('Monthly Operating Costs')}!{CL(GC+m-1)}${r_tot}"

# =====================================================================
# RETAIL READINESS COSTS
# =====================================================================
wr = ws_["Retail Readiness Costs"]
title(wr, "Retail readiness costs: Wave 3 board-book launch and big-box readiness (all [VERIFY])",
      "Almost nothing here has a live quote yet. Get 2-3 written quotes for each line before any commitment (BLIND-SPOTS #11). "
      "Section A hits in the board-book print month; Section B hits 'lead time' months before retail launch; Section C runs monthly while retail is live.")
wr.merge_cells("A2:K3"); wr["A2"].alignment = WRAP; wr.row_dimensions[2].height = 28
hdr(wr, 5, 1, 11, ["Item", "Section", "Low ($)", "High ($)", "Model ($ per unit of qty)", "Basis", "Qty", "Model total ($)",
                   "Timing", "Source", "Status"], 32)
RR_A = [
    ("CPSIA third-party lab testing, board book for ages 0-3", "A", 300, 1000, "per SKU", 1, "Print month", "PROTECTION-PLAN #9: 'a few hundred dollars per product'; books for 3 and under are not exempt.", "Repo [VERIFY]"),
    ("Children's Product Certificate (prepared in-house)", "A", 0, 0, "per SKU", 1, "Print month", "PROTECTION-PLAN: CPC $0 to prepare.", "Repo"),
    ("Tracking labels and printer file changes", "A", 0, 200, "one-time", 1, "Print month", "Not in the repo.", "[VERIFY]"),
    ("Board-book printer hard proof and quote samples", "A", 100, 400, "one-time", 1, "Print month", "Not in the repo; BLIND-SPOTS #11 asks for 2-3 full quotes.", "[VERIFY]"),
    ("3PL onboarding / setup", "A", 0, 500, "one-time", 1, "Print month", "Not in the repo.", "[VERIFY]"),
    ("Amazon FBA inbound prep and labels", "A", 50, 200, "one-time", 1, "Print month", "Not in the repo.", "[VERIFY]"),
    ("Attorney check of pre-sale terms (FTC mail-order rule)", "A", 300, 700, "one-time", 1, "Print month", "BLIND-SPOTS #16-#17: lawyer check $300-$700.", "Repo [VERIFY]"),
]
RR_B = [
    ("GS1 US company prefix for non-book SKUs (card deck, boxed sets); books use ISBN barcodes", "B", 250, 750, "one-time", 1, "Retail launch - lead", "AMAZON-AND-RETAIL-ROADMAP B5 requires GS1 barcodes; fee not in the repo.", "[VERIFY]"),
    ("CPSIA testing for more children's SKUs (card deck, 3-pack box)", "B", 300, 1000, "per SKU", 2, "Retail launch - lead", "PROTECTION-PLAN; ROADMAP B5.", "Repo [VERIFY]"),
    ("Retail-ready packaging design and dielines", "B", 500, 2000, "one-time", 1, "Retail launch - lead", "ROADMAP B5 'retail-ready packaging'; no figure in the repo.", "[VERIFY]"),
    ("Attorney review of retailer / distributor vendor agreement", "B", 500, 2500, "one-time", 1, "Retail launch - lead", "Repo attorney range $300-$2,500 per document (PROTECTION-PLAN).", "Repo [VERIFY]"),
    ("Buyer sell sheet, line list and sample kits", "B", 200, 600, "one-time", 1, "Retail launch - lead", "ROADMAP B7 (pitch only with proof); no figure in the repo.", "[VERIFY]"),
    ("Faire brand application and first-order sample stock allowance", "B", 0, 300, "one-time", 1, "Retail launch - lead", "storefront-setup-guide §19.", "[VERIFY]"),
]
RR_C = [
    ("3PL EDI connection and retailer compliance", "C", 100, 300, "per month", 1, "Monthly while retail is live", "ROADMAP B5: 'EDI capability through a 3PL'; no figure in the repo.", "[VERIFY]"),
    ("Insurance uplift to retailer limits ($1M+ occurrence, umbrella, additional insureds)", "C", 50, 250, "per month", 1, "Monthly while retail is live", "PROTECTION-PLAN retailer requirements (Amazon, Faire+); premium uplift not in the repo.", "[VERIFY]"),
    ("GS1 annual renewal (monthly equivalent)", "C", 4, 13, "per month", 1, "Monthly while retail is live", "Not in the repo.", "[VERIFY]"),
]
RR_ROWS = {}
r = 6
def rr_block(items, label):
    global r
    sec(wr, r, 1, 11, label); r += 1
    first = r
    for (item, s, lo, hi, basis, qty, timing, src, st) in items:
        put(wr, f"A{r}", item); put(wr, f"B{r}", s)
        put(wr, f"C{r}", lo, fmt=CUR0, font=F_IN); put(wr, f"D{r}", hi, fmt=CUR0, font=F_IN)
        put(wr, f"E{r}", f"=C{r}+{A['pos']}*(D{r}-C{r})", fmt=CUR0)
        put(wr, f"F{r}", basis); put(wr, f"G{r}", qty, fmt="0", font=F_IN)
        put(wr, f"H{r}", f"=E{r}*G{r}", fmt=CUR0)
        put(wr, f"I{r}", timing); put(wr, f"J{r}", src); put(wr, f"K{r}", st)
        r += 1
    last = r - 1
    put(wr, f"A{r}", f"Total, section {items[0][1]}", bold=True)
    put(wr, f"H{r}", f"=SUM(H{first}:H{last})", fmt=CUR0, bold=True)
    for c in range(1, 12):
        wr.cell(row=r, column=c).fill = TOT_FILL
    tot = r
    r += 2
    return tot
tA = rr_block(RR_A, "A. Wave 3 board-book launch (one-time, in the print month)")
tB = rr_block(RR_B, "B. Big-box and wholesale readiness (one-time, before retail launch)")
tC = rr_block(RR_C, "C. Ongoing while retail is live (per month)")
RR = dict(A=f"{q('Retail Readiness Costs')}!$H${tA}", B=f"{q('Retail Readiness Costs')}!$H${tB}", C=f"{q('Retail Readiness Costs')}!$H${tC}")
sec(wr, r, 1, 11, "D. Inventory: first offset run (illustrative at the minimum run; the forecast sizes each scenario's run from pre-sale orders)"); r += 1
put(wr, f"A{r}", "Offset print cost per copy (from Assumptions)"); put(wr, f"E{r}", f"={A['bb_print']}", fmt=CUR2); r += 1
put(wr, f"A{r}", "Landed cost per copy incl. freight/duties"); put(wr, f"E{r}", f"={A['landed']}", fmt=CUR2); r += 1
put(wr, f"A{r}", "Minimum run (copies)"); put(wr, f"G{r}", f"={A['min_run']}", fmt=NUM0); r += 1
put(wr, f"A{r}", "Cash for a minimum first run", bold=True); put(wr, f"H{r}", f"=E{r-2}*G{r-1}", fmt=CUR0, bold=True)
RR["run_illus"] = f"{q('Retail Readiness Costs')}!$H${r}"; r += 1
put(wr, f"A{r}", "Landed cost as % of retail price (target at or below 35-40%, DEMAND-CHECK rule 9)")
put(wr, f"H{r}", f"=E{r-3}/{A['bb_price']}", fmt=PCT1); r += 2
sec(wr, r, 1, 11, "E. Ongoing Wave 3 fixed costs (from Assumptions; charged monthly once the first run is printed)"); r += 1
put(wr, f"A{r}", "3PL monthly minimum / storage"); put(wr, f"H{r}", f"={A['tpl_min']}", fmt=CUR0); r += 1
put(wr, f"A{r}", "Amazon Seller Central Professional plan"); put(wr, f"H{r}", f"={A['seller_central']}", fmt=CUR2); r += 2
put(wr, f"A{r}", "Readiness items the model does not price: a distributor or sales rep (commission), retailer co-op or placement fees, and in-store "
                 "(brick-and-mortar) Target or Walmart vendor setup. Books usually reach big-box shelves through book distributors or a publishing deal "
                 "(AMAZON-AND-RETAIL-ROADMAP B6) [VERIFY].", font=F_NOTE)
wr.merge_cells(f"A{r}:K{r+1}"); wr[f"A{r}"].alignment = WRAP; wr.row_dimensions[r].height = 30
for c, w in zip("ABCDEFGHIJK", [70, 8, 10, 10, 12, 10, 6, 12, 22, 70, 14]):
    wr.column_dimensions[c].width = w
wr.freeze_panes = "B6"

# =====================================================================
# 36-MONTH FORECAST
# =====================================================================
wf = ws_["36-Month Forecast"]
FS = "36-Month Forecast"
title(wf, "36-month forecast: three scenarios, driven by traffic -> conversion -> orders -> order value, by channel and wave",
      "Traffic = trend (launch traffic compounded by monthly growth) x seasonality. Conversion ramps up over the review-ramp months. "
      "Net contribution = items x blended net per item from Unit Economics. Figures are planning estimates, not promises of income.")
wf.merge_cells("A2:P3"); wf["A2"].alignment = WRAP; wf.row_dimensions[2].height = 28
FC = 3
def mc(m): return CL(FC + m - 1)
YC = {1: CL(FC + 36), 2: CL(FC + 37), 3: CL(FC + 38), "T": CL(FC + 39)}
put(wf, "A4", "Model month #", bold=True); put(wf, "A5", "Month", bold=True); put(wf, "A6", "Model year", bold=True)
put(wf, "A7", "Seasonality: consumer", bold=True); put(wf, "A8", "Seasonality: school / group", bold=True)
for m in range(1, 37):
    c = mc(m)
    put(wf, f"{c}4", m, fmt="0")
    put(wf, f"{c}5", f"=DATE(YEAR({A['start']}),MONTH({A['start']})+{c}4-1,1)", fmt=DATEF)
    put(wf, f"{c}6", f"=INT(({c}4-1)/12)+1", fmt="0")
    put(wf, f"{c}7", f"=INDEX({A['season_c']},MONTH({c}5))", fmt=NUM2)
    put(wf, f"{c}8", f"=INDEX({A['season_s']},MONTH({c}5))", fmt=NUM2)
hdr(wf, 5, FC, FC + 39)
for m in range(1, 37):
    wf[f"{mc(m)}5"].number_format = DATEF
for k, lab in ((1, "Year 1"), (2, "Year 2"), (3, "Year 3"), ("T", "36 months")):
    wf[f"{YC[k]}5"].value = lab
    wf[f"{YC[k]}4"].value = "Oct-Sep" if k != "T" else "Total"
    wf[f"{YC[k]}4"].font = F_NOTE
put(wf, "B4", "Unit / note", font=F_NOTE)

FR = {}   # FR[s][key] -> row
BLOCK_START = 10

def year_cols(ws, row, kind, fmt):
    """kind: 'sum' | 'end' | None"""
    first, last = mc(1), mc(36)
    for k in (1, 2, 3, "T"):
        c = YC[k]
        if kind == "sum":
            if k == "T":
                f = f"=SUM({first}{row}:{last}{row})"
            else:
                f = f"=SUMIF(${first}$6:${last}$6,{k},{first}{row}:{last}{row})"
        elif kind == "end":
            mm = {1: 12, 2: 24, 3: 36, "T": 36}[k]
            f = f"={mc(mm)}{row}"
        else:
            continue
        put(ws, f"{c}{row}", f, fmt=fmt)
        ws[f"{c}{row}"].fill = TOT_FILL

def build_block(s, start):
    rows = {}
    keys = ["hdr", "print", "presale", "run", "need", "blank0"]
    for code, *_ in CH:
        keys += [f"{code}_hdr", f"{code}_act", f"{code}_trend", f"{code}_traffic", f"{code}_cvr", f"{code}_orders",
                 f"{code}_units", f"{code}_aov", f"{code}_gross", f"{code}_contrib"]
    keys += ["email_hdr", "email_new", "email_churn", "email_end",
             "tot_hdr", "orders", "units", "cum_orders", "cum_units", "gross", "contrib", "margin",
             "ads", "opex", "wavefix", "netop", "startup", "waveone", "netafter", "cumnet",
             "inv_hdr", "phys", "cumphys", "copies", "inv", "cumcopies", "onhand",
             "flag_hdr", "flag_first", "flag_sust"]
    for i, k in enumerate(keys):
        rows[k] = start + i
    FR[s] = rows
    S = lambda key: A[key][s]
    rr = rows
    # header
    hdr(wf, rr["hdr"], 1, FC + 39, [f"{SCN[s].upper()} SCENARIO"])
    # scalars
    put(wf, f"A{rr['print']}", "Board-book print month (pre-sale ends), model month #")
    put(wf, f"B{rr['print']}", f"=IF({S('BOARD_inc')}=1,{S('BOARD_launch')}+{A['presale_len']},0)", fmt="0")
    put(wf, f"A{rr['presale']}", "Pre-sale units ordered before the print month (copies)")
    u = rr["BOARD_units"]
    put(wf, f"B{rr['presale']}", f"=IF($B${rr['print']}=0,0,SUMPRODUCT(($C$4:${mc(36)}$4>={S('BOARD_launch')})*($C$4:${mc(36)}$4<$B${rr['print']})*($C${u}:${mc(36)}${u})))", fmt=NUM0)
    put(wf, f"A{rr['run']}", "First offset run size = max(minimum run, pre-sale units x (1 + buffer)), rounded up to 100")
    put(wf, f"B{rr['run']}", f"=IF($B${rr['print']}=0,0,MAX({A['min_run']},ROUNDUP($B${rr['presale']}*(1+{A['buffer']})/100,0)*100))", fmt=NUM0)
    put(wf, f"A{rr['need']}", "Pre-sale copies needed to fund a minimum run + launch costs (go / no-go line; BLIND-SPOTS #16)")
    put(wf, f"B{rr['need']}", f"=IF($B${rr['print']}=0,0,ROUNDUP(({A['min_run']}*{A['landed']}+{RR['A']})*(1+{A['presale_buf']})/({A['bb_price']}*(1-{A['shop_pct']})-{A['shop_fix']}-{A['pick']}),0))", fmt=NUM0)
    for k in ("print", "presale", "run", "need"):
        wf[f"B{rr[k]}"].fill = TOT_FILL
    for code, name, tl, lag, seas, src in CH:
        sec(wf, rr[f"{code}_hdr"], 1, FC + 39, name)
        labels = {
            "act": ("Active (1 = live)", "flag"), "trend": ("Traffic trend (before seasonality)", "visits"),
            "traffic": (tl, "visits"), "cvr": ("Conversion rate (with review ramp)", "%"),
            "orders": ("Orders", "orders"), "units": ("Items sold", "items"), "aov": ("Average order value", "$"),
            "gross": ("Gross sales (price paid by customers)", "$"), "contrib": ("Net contribution to AlphaPlay", "$"),
        }
        if code == "COURSE":
            labels["trend"] = ("Subscribers reached (list last month x reach)", "people")
        if code == "SITE":
            labels["orders"] = ("Orders (US sessions; international go to the MoR once it is live)", "orders")
        for k, (lab, unit) in labels.items():
            put(wf, f"A{rr[code+'_'+k]}", "   " + lab)
            put(wf, f"B{rr[code+'_'+k]}", unit, font=F_NOTE)
        seas_row = 7 if seas == "c" else 8
        for m in range(1, 37):
            c, p = mc(m), mc(m - 1) if m > 1 else None
            ra = rr[f"{code}_act"]
            put(wf, f"{c}{ra}", f"=IF(AND({S(code+'_inc')}=1,{c}$4>={S(code+'_launch')}),1,0)", fmt="0")
            rt_ = rr[f"{code}_trend"]
            if code == "MOR":
                f = f"=IF({c}{ra}=0,0,{c}{rr['SITE_trend']}*{S('intl')})"
            elif code == "COURSE":
                f = f"=IF({c}{ra}=0,0,{('0' if m == 1 else p + str(rr['email_end']))}*{S('reach')})"
            else:
                g = f"CHOOSE({c}$6,{S(code+'_g1')},{S(code+'_g2')},{S(code+'_g3')})"
                if m == 1:
                    f = f"=IF({c}{ra}=0,0,{S(code+'_base')})"
                else:
                    f = f"=IF({c}{ra}=0,0,IF({c}$4={S(code+'_launch')},{S(code+'_base')},{p}{rt_}*(1+{g})))"
            put(wf, f"{c}{rt_}", f, fmt=NUM0)
            put(wf, f"{c}{rr[code+'_traffic']}", f"={c}{rt_}*{c}${seas_row}", fmt=NUM0)
            put(wf, f"{c}{rr[code+'_cvr']}", f"={S(code+'_cvr')}*MIN(1,({c}$4-{S(code+'_launch')}+1)/{A['ramp']})*{c}{ra}", fmt=PCT2)
            if code == "SITE":
                f = f"=({c}{rr['SITE_traffic']}-{c}{rr['MOR_traffic']})*{c}{rr['SITE_cvr']}"
            else:
                f = f"={c}{rr[code+'_traffic']}*{c}{rr[code+'_cvr']}"
            put(wf, f"{c}{rr[code+'_orders']}", f, fmt=NUM1)
            put(wf, f"{c}{rr[code+'_units']}", f"={c}{rr[code+'_orders']}*{S(code+'_items')}", fmt=NUM1)
            put(wf, f"{c}{rr[code+'_aov']}", f"={UE[code]['price']}*{S(code+'_items')}*{c}{ra}", fmt=CUR2)
            put(wf, f"{c}{rr[code+'_gross']}", f"={c}{rr[code+'_units']}*{UE[code]['price']}", fmt=CUR0)
            put(wf, f"{c}{rr[code+'_contrib']}", f"={c}{rr[code+'_units']}*{UE[code]['net']}", fmt=CUR0)
        for k, kind, fmt in (("traffic", "sum", NUM0), ("orders", "sum", NUM0), ("units", "sum", NUM0),
                             ("gross", "sum", CUR0), ("contrib", "sum", CUR0)):
            year_cols(wf, rr[f"{code}_{k}"], kind, fmt)
    # email
    sec(wf, rr["email_hdr"], 1, FC + 39, "Email list (owned audience)")
    for k, lab in (("email_new", "   New subscribers (site visitors + buyers via bonus QR)"), ("email_churn", "   Unsubscribes"),
                   ("email_end", "   Subscribers at month end")):
        put(wf, f"A{rr[k]}", lab); put(wf, f"B{rr[k]}", "people", font=F_NOTE)
    buyer_codes = ["SITE", "ETSY", "KDP", "INGRAM", "MOR", "BOARD"]
    for m in range(1, 37):
        c, p = mc(m), (mc(m - 1) if m > 1 else None)
        orders_sum = "+".join(f"{c}{rr[x+'_orders']}" for x in buyer_codes)
        put(wf, f"{c}{rr['email_new']}", f"={c}{rr['SITE_traffic']}*{S('signup')}+({orders_sum})*{S('optin')}", fmt=NUM0)
        put(wf, f"{c}{rr['email_churn']}", "=0" if m == 1 else f"={p}{rr['email_end']}*{S('churn')}", fmt=NUM0)
        put(wf, f"{c}{rr['email_end']}", f"={c}{rr['email_new']}-{c}{rr['email_churn']}" + ("" if m == 1 else f"+{p}{rr['email_end']}"), fmt=NUM0)
    year_cols(wf, rr["email_new"], "sum", NUM0); year_cols(wf, rr["email_churn"], "sum", NUM0); year_cols(wf, rr["email_end"], "end", NUM0)
    # totals
    sec(wf, rr["tot_hdr"], 1, FC + 39, "Totals and operating result")
    TL = {
        "orders": ("Total orders", "orders", NUM0, "sum"), "units": ("Total items sold", "items", NUM0, "sum"),
        "cum_orders": ("Cumulative orders", "orders", NUM0, "end"), "cum_units": ("Cumulative items sold", "items", NUM0, "end"),
        "gross": ("Total gross sales (customer prices, excl. sales tax/VAT)", "$", CUR0, "sum"),
        "contrib": ("Total net contribution (after platform, payment, print/POD and fulfilment costs)", "$", CUR0, "sum"),
        "margin": ("Net contribution as % of gross sales", "%", PCT1, None),
        "ads": ("Paid advertising", "$", CUR0, "sum"), "opex": ("Operating costs (Monthly Operating Costs tab)", "$", CUR0, "sum"),
        "wavefix": ("Wave fixed costs (3PL minimum, Seller Central, retail EDI/insurance uplift)", "$", CUR0, "sum"),
        "netop": ("NET OPERATING RESULT (before one-time costs)", "$", CUR0, "sum"),
        "startup": ("One-time startup costs (Startup Costs tab)", "$", CUR0, "sum"),
        "waveone": ("One-time wave launch and retail readiness costs", "$", CUR0, "sum"),
        "netafter": ("Net result after one-time costs", "$", CUR0, "sum"),
        "cumnet": ("Cumulative net result", "$", CUR0, "end"),
        "phys": ("   Physical items sold from held stock (board book + retail)", "items", NUM0, "sum"),
        "cumphys": ("   Cumulative physical items", "items", NUM0, "end"),
        "copies": ("   Copies ordered from the printer (first run + reorders)", "copies", NUM0, "sum"),
        "inv": ("   Inventory purchases (cash; not in operating result)", "$", CUR0, "sum"),
        "cumcopies": ("   Cumulative copies ordered", "copies", NUM0, "end"),
        "onhand": ("   Copies at the 3PL (negative = pre-sale copies still owed)", "copies", NUM0, "end"),
        "flag_first": ("   Month # if operating result >= 0", "month", "0", None),
        "flag_sust": ("   Month # if operating result stays >= 0 through month 36", "month", "0", None),
    }
    sec(wf, rr["inv_hdr"], 1, FC + 39, "Held inventory (Wave 3 board book and retail; stock sits at a 3PL or Amazon, never with the founder)")
    sec(wf, rr["flag_hdr"], 1, FC + 39, "Break-even helper rows")
    for k, (lab, unit, fmt, kind) in TL.items():
        put(wf, f"A{rr[k]}", lab, bold=k in ("netop", "gross", "contrib"))
        put(wf, f"B{rr[k]}", unit, font=F_NOTE)
    codes = [c[0] for c in CH]
    for m in range(1, 37):
        c, p = mc(m), (mc(m - 1) if m > 1 else None)
        put(wf, f"{c}{rr['orders']}", "=" + "+".join(f"{c}{rr[x+'_orders']}" for x in codes), fmt=NUM0)
        put(wf, f"{c}{rr['units']}", "=" + "+".join(f"{c}{rr[x+'_units']}" for x in codes), fmt=NUM0)
        put(wf, f"{c}{rr['cum_orders']}", f"={c}{rr['orders']}" + ("" if m == 1 else f"+{p}{rr['cum_orders']}"), fmt=NUM0)
        put(wf, f"{c}{rr['cum_units']}", f"={c}{rr['units']}" + ("" if m == 1 else f"+{p}{rr['cum_units']}"), fmt=NUM0)
        put(wf, f"{c}{rr['gross']}", "=" + "+".join(f"{c}{rr[x+'_gross']}" for x in codes), fmt=CUR0)
        put(wf, f"{c}{rr['contrib']}", "=" + "+".join(f"{c}{rr[x+'_contrib']}" for x in codes), fmt=CUR0)
        put(wf, f"{c}{rr['margin']}", f"=IF({c}{rr['gross']}=0,0,{c}{rr['contrib']}/{c}{rr['gross']})", fmt=PCT1)
        put(wf, f"{c}{rr['ads']}", f"=IF({c}$4>={S('ads_start')},CHOOSE({c}$6,{S('ads1')},{S('ads2')},{S('ads3')}),0)", fmt=CUR0)
        put(wf, f"{c}{rr['opex']}", f"={OPTOT(m)}", fmt=CUR0)
        put(wf, f"{c}{rr['wavefix']}",
            f"=IF(AND({S('BOARD_inc')}=1,$B${rr['print']}>0,{c}$4>=$B${rr['print']}),{A['tpl_min']}+{A['seller_central']},0)+IF({c}{rr['RETAIL_act']}=1,{RR['C']},0)", fmt=CUR0)
        put(wf, f"{c}{rr['netop']}", f"={c}{rr['contrib']}-{c}{rr['ads']}-{c}{rr['opex']}-{c}{rr['wavefix']}", fmt=CUR0)
        put(wf, f"{c}{rr['startup']}", f"=SUMPRODUCT(({SU_RANGE['F']}={c}$4)*{SU_RANGE['G']}*{SU_RANGE['E']})", fmt=CUR0)
        put(wf, f"{c}{rr['waveone']}",
            f"=IF(AND({S('BOARD_inc')}=1,{c}$4=$B${rr['print']}),{RR['A']},0)"
            f"+IF(AND({S('RETAIL_inc')}=1,{c}$4=MAX(1,{S('RETAIL_launch')}-{A['retail_lead']})),{RR['B']},0)"
            f"+IF(AND({S('SCHOOL_inc')}=1,{c}$4={S('SCHOOL_launch')}),{A['school_setup']},0)", fmt=CUR0)
        put(wf, f"{c}{rr['netafter']}", f"={c}{rr['netop']}-{c}{rr['startup']}-{c}{rr['waveone']}", fmt=CUR0)
        put(wf, f"{c}{rr['cumnet']}", f"={c}{rr['netafter']}" + ("" if m == 1 else f"+{p}{rr['cumnet']}"), fmt=CUR0)
        put(wf, f"{c}{rr['phys']}", f"={c}{rr['BOARD_units']}+{c}{rr['RETAIL_units']}", fmt=NUM0)
        put(wf, f"{c}{rr['cumphys']}", f"={c}{rr['phys']}" + ("" if m == 1 else f"+{p}{rr['cumphys']}"), fmt=NUM0)
        run = f"$B${rr['run']}"; pm = f"$B${rr['print']}"
        prevcum = "0" if m == 1 else f"{p}{rr['cumphys']}"
        put(wf, f"{c}{rr['copies']}",
            f"=IF({run}=0,0,IF({c}$4={pm},{run},IF({c}$4>{pm},(INT({c}{rr['cumphys']}/{run})-INT({prevcum}/{run}))*{run},0)))", fmt=NUM0)
        put(wf, f"{c}{rr['inv']}", f"={c}{rr['copies']}*{A['landed']}", fmt=CUR0)
        put(wf, f"{c}{rr['cumcopies']}", f"={c}{rr['copies']}" + ("" if m == 1 else f"+{p}{rr['cumcopies']}"), fmt=NUM0)
        put(wf, f"{c}{rr['onhand']}", f"={c}{rr['cumcopies']}-{c}{rr['cumphys']}", fmt=NUM0)
        put(wf, f"{c}{rr['flag_first']}", f"=IF({c}{rr['netop']}>=0,{c}$4,\"\")", fmt="0")
        put(wf, f"{c}{rr['flag_sust']}", f"=IF(COUNTIF({c}{rr['netop']}:${mc(36)}{rr['netop']},\"<0\")=0,{c}$4,\"\")", fmt="0")
    for k, (lab, unit, fmt, kind) in TL.items():
        if kind:
            year_cols(wf, rr[k], kind, fmt)
    for k in ("netop", "gross", "contrib"):
        for cc in range(1, FC + 40):
            wf.cell(row=rr[k], column=cc).font = Font(name=AR, bold=True, color=wf.cell(row=rr[k], column=cc).font.color)
            wf.cell(row=rr[k], column=cc).border = TOP
    return start + len(keys) + 1

nxt = BLOCK_START
for s in range(3):
    nxt = build_block(s, nxt)
wf.column_dimensions["A"].width = 66
wf.column_dimensions["B"].width = 11
for m in range(1, 37):
    wf.column_dimensions[mc(m)].width = 10
for k in YC.values():
    wf.column_dimensions[k].width = 12
wf.freeze_panes = "C9"

# =====================================================================
# CASH FLOW
# =====================================================================
wc = ws_["Cash Flow"]
title(wc, "Cash flow: payout delays, inventory purchases, one-time costs and the tax reserve",
      "Marketplace payouts arrive after the sale month (KDP about 60 days, IngramSpark about 90 days; see Assumptions section 4). "
      "Costs are negative. Sales tax and VAT are pass-through and excluded. The tax reserve is money set aside, not tax actually due; the accountant sets the real figure.")
wc.merge_cells("A2:P3"); wc["A2"].alignment = WRAP; wc.row_dimensions[2].height = 28
put(wc, "A4", "Model month #", bold=True); put(wc, "A5", "Month", bold=True); put(wc, "A6", "Model year", bold=True)
for m in range(1, 37):
    c = mc(m)
    put(wc, f"{c}4", f"={q(FS)}!{c}4", fmt="0")
    put(wc, f"{c}5", f"={q(FS)}!{c}5", fmt=DATEF)
    put(wc, f"{c}6", f"={q(FS)}!{c}6", fmt="0")
hdr(wc, 5, FC, FC + 39)
for m in range(1, 37):
    wc[f"{mc(m)}5"].number_format = DATEF
for k, lab in ((1, "Year 1"), (2, "Year 2"), (3, "Year 3"), ("T", "36 months")):
    wc[f"{YC[k]}5"].value = lab
CFR = {}
r = 8
for s in range(3):
    rr = {}
    keys = ["hdr"] + [f"pay_{c[0]}" for c in CH] + ["pay_tot", "owner", "inv", "ads", "opex", "wavefix", "startup", "waveone",
                                                     "net", "resv", "opbal", "resbal", "total", "target", "gap", "flag_cash"]
    for i, k in enumerate(keys):
        rr[k] = r + i
    CFR[s] = rr
    fr = FR[s]
    hdr(wc, rr["hdr"], 1, FC + 39, [f"{SCN[s].upper()} SCENARIO"])
    lab = {"pay_tot": "Total payouts received", "owner": "Owner capital contribution",
           "inv": "Inventory purchases (offset print runs, landed)", "ads": "Paid advertising", "opex": "Operating costs",
           "wavefix": "Wave fixed costs", "startup": "One-time startup costs", "waveone": "One-time wave launch / retail readiness",
           "net": "NET CASH FLOW before tax reserve", "resv": "Transfer to tax reserve account",
           "opbal": "OPERATING ACCOUNT BALANCE (month end)", "resbal": "Tax reserve account balance", "total": "Total cash (both accounts)",
           "target": "Reserve target (months x fixed costs incl. ads)", "gap": "Operating balance above / (below) reserve target",
           "flag_cash": "   Helper: month # if operating balance stays >= 0 through month 36"}
    for code, name, *_ in CH:
        lab[f"pay_{code}"] = f"   Payout: {name}"
    for k, t in lab.items():
        put(wc, f"A{rr[k]}", t, bold=k in ("net", "opbal", "pay_tot"))
    for m in range(1, 37):
        c, p = mc(m), (mc(m - 1) if m > 1 else None)
        for code, *_ in CH:
            lag = A[f"lag_{code}"]
            cr = fr[f"{code}_contrib"]; ur = fr[f"{code}_units"]
            put(wc, f"{c}{rr['pay_'+code]}",
                f"=IF({c}$4-{lag}>=1,INDEX({q(FS)}!$C${cr}:${mc(36)}${cr},1,{c}$4-{lag})+INDEX({q(FS)}!$C${ur}:${mc(36)}${ur},1,{c}$4-{lag})*{UE[code]['landed']},0)",
                fmt=CUR0)
        put(wc, f"{c}{rr['pay_tot']}", f"=SUM({c}{rr['pay_SITE']}:{c}{rr['pay_RETAIL']})", fmt=CUR0)
        put(wc, f"{c}{rr['owner']}", f"={A['owner']}" if m == 1 else "=0", fmt=CUR0)
        for k in ("inv", "ads", "opex", "wavefix", "startup", "waveone"):
            put(wc, f"{c}{rr[k]}", f"=-{q(FS)}!{c}{fr[k]}", fmt=CUR0)
        put(wc, f"{c}{rr['net']}", f"=SUM({c}{rr['pay_tot']}:{c}{rr['waveone']})", fmt=CUR0)
        put(wc, f"{c}{rr['resv']}", f"=-{A['taxres']}*MAX(0,{c}{rr['net']})", fmt=CUR0)
        prev_op = A["open"] if m == 1 else f"{p}{rr['opbal']}"
        put(wc, f"{c}{rr['opbal']}", f"={prev_op}+{c}{rr['net']}+{c}{rr['resv']}", fmt=CUR0)
        put(wc, f"{c}{rr['resbal']}", f"=-{c}{rr['resv']}" + ("" if m == 1 else f"+{p}{rr['resbal']}"), fmt=CUR0)
        put(wc, f"{c}{rr['total']}", f"={c}{rr['opbal']}+{c}{rr['resbal']}", fmt=CUR0)
        put(wc, f"{c}{rr['target']}", f"={A['resmo']}*({q(FS)}!{c}{fr['opex']}+{q(FS)}!{c}{fr['wavefix']}+{q(FS)}!{c}{fr['ads']})", fmt=CUR0)
        put(wc, f"{c}{rr['gap']}", f"={c}{rr['opbal']}-{c}{rr['target']}", fmt=CUR0)
        put(wc, f"{c}{rr['flag_cash']}", f"=IF(COUNTIF({c}{rr['opbal']}:${mc(36)}{rr['opbal']},\"<0\")=0,{c}$4,\"\")", fmt="0")
    for k in [f"pay_{c[0]}" for c in CH] + ["pay_tot", "owner", "inv", "ads", "opex", "wavefix", "startup", "waveone", "net", "resv"]:
        year_cols(wc, rr[k], "sum", CUR0)
    for k in ("opbal", "resbal", "total", "target", "gap"):
        year_cols(wc, rr[k], "end", CUR0)
    for k in ("net", "opbal"):
        for cc in range(1, FC + 40):
            cell = wc.cell(row=rr[k], column=cc)
            cell.font = Font(name=AR, bold=True, color=cell.font.color); cell.border = TOP
    r = r + len(keys) + 1
wc.column_dimensions["A"].width = 62
wc.column_dimensions["B"].width = 4
for m in range(1, 37):
    wc.column_dimensions[mc(m)].width = 10
for k in YC.values():
    wc.column_dimensions[k].width = 12
wc.freeze_panes = "C7"

# =====================================================================
# BREAK-EVEN
# =====================================================================
wb_ = ws_["Break-even"]
title(wb_, "Break-even: when, and how many orders or units",
      "Operating break-even excludes one-time startup, wave-launch and inventory costs. Cash payback includes everything, including payout delays. "
      "'Sustained' means the result stays at or above zero in every later month through month 36.")
wb_.merge_cells("A2:E3"); wb_["A2"].alignment = WRAP; wb_.row_dimensions[2].height = 28
hdr(wb_, 5, 1, 5, ["Measure", "Conservative", "Expected", "Strong", "Note"], 22)
BE = {}
def dt(x):
    return f"=IF(ISNUMBER({x}),DATE(YEAR({A['start']}),MONTH({A['start']})+{x}-1,1),\"Not within 36 months\")"
rows_be = []
r = 6
def be_row(key, label, fn, fmt, note=""):
    global r
    put(wb_, f"A{r}", label)
    for s in range(3):
        put(wb_, 'BCD'[s] + str(r), fn(s), fmt=fmt)
    put(wb_, f"E{r}", note, font=F_NOTE)
    BE[key] = r
    r += 1
F36 = mc(36)
def fref(s, key, col=None):
    return f"{q(FS)}!${mc(1)}${FR[s][key]}:${F36}${FR[s][key]}"
sec(wb_, r, 1, 5, "When"); r += 1
be_row("first_m", "First month with operating result >= 0 (month #)",
       lambda s: f"=IF(COUNT({fref(s,'flag_first')})=0,\"Not reached\",MIN({fref(s,'flag_first')}))", "0")
be_row("first_d", "   ...calendar month", lambda s: dt(f"{'BCD'[s]}{BE['first_m']}"), DATEF)
be_row("sust_m", "Sustained operating break-even (month #)",
       lambda s: f"=IF(COUNT({fref(s,'flag_sust')})=0,\"Not reached\",MIN({fref(s,'flag_sust')}))", "0",
       "Headline break-even measure used on the Dashboard.")
be_row("sust_d", "   ...calendar month", lambda s: dt(f"{'BCD'[s]}{BE['sust_m']}"), DATEF)
def cfref(s, key):
    return f"{q('Cash Flow')}!${mc(1)}${CFR[s][key]}:${F36}${CFR[s][key]}"
be_row("cash_m", "Cash payback: operating balance >= 0 from this month on (month #)",
       lambda s: f"=IF(COUNT({cfref(s,'flag_cash')})=0,\"Not reached\",MIN({cfref(s,'flag_cash')}))", "0",
       "Includes payout delays, startup costs, inventory and the tax-reserve transfers.")
be_row("cash_d", "   ...calendar month", lambda s: dt(f"{'BCD'[s]}{BE['cash_m']}"), DATEF)
be_row("peak", "Peak funding need (lowest operating balance, shown positive)",
       lambda s: f"=MAX(0,-MIN({cfref(s,'opbal')}))", CUR0, "Cash the founder must put in (or not spend) to avoid a negative balance.")
be_row("peak_m", "   ...month # of the lowest balance",
       lambda s: f"=MATCH(MIN({cfref(s,'opbal')}),{cfref(s,'opbal')},0)", "0")
r += 1
sec(wb_, r, 1, 5, "How many (volume at break-even)"); r += 1
def at_month(s, key, mcell):
    return f"=IF(ISNUMBER({mcell}),INDEX({fref(s,key)},1,{mcell}),\"-\")"
be_row("cum_orders_be", "Cumulative orders by the sustained break-even month",
       lambda s: at_month(s, "cum_orders", f"{'BCD'[s]}{BE['sust_m']}"), NUM0)
be_row("cum_units_be", "Cumulative items sold by the sustained break-even month",
       lambda s: at_month(s, "cum_units", f"{'BCD'[s]}{BE['sust_m']}"), NUM0)
be_row("orders_be", "Orders in the sustained break-even month",
       lambda s: at_month(s, "orders", f"{'BCD'[s]}{BE['sust_m']}"), NUM0)
r += 1
sec(wb_, r, 1, 5, "Orders needed each month to cover fixed costs (operating costs + wave fixed costs + ads)"); r += 1
for mm in (6, 12, 24, 36):
    c = mc(mm)
    be_row(f"fixed_{mm}", f"Month {mm}: fixed costs",
           lambda s, c=c: f"={q(FS)}!{c}{FR[s]['opex']}+{q(FS)}!{c}{FR[s]['wavefix']}+{q(FS)}!{c}{FR[s]['ads']}", CUR0)
    be_row(f"cpo_{mm}", f"Month {mm}: net contribution per order",
           lambda s, c=c: f"=IF({q(FS)}!{c}{FR[s]['orders']}=0,0,{q(FS)}!{c}{FR[s]['contrib']}/{q(FS)}!{c}{FR[s]['orders']})", CUR2)
    be_row(f"need_{mm}", f"Month {mm}: orders needed to cover fixed costs",
           lambda s, mm=mm: f"=IF({'BCD'[s]}{BE[f'cpo_{mm}']}<=0,\"-\",{'BCD'[s]}{BE[f'fixed_{mm}']}/{'BCD'[s]}{BE[f'cpo_{mm}']})", NUM0)
    be_row(f"act_{mm}", f"Month {mm}: orders the scenario forecasts",
           lambda s, c=c: f"={q(FS)}!{c}{FR[s]['orders']}", NUM0)
r += 1
sec(wb_, r, 1, 5, "Single-product view: units per month of ONE product that would cover the Expected month-12 fixed costs"); r += 1
hdr(wb_, r, 1, 5, ["Product and channel", "Net per unit ($)", "Units per month needed", "Units per day", "Source row"]); r += 1
fixed12 = f"$C${BE['fixed_12']}"
for pid in ("ETSY_routine", "ETSY_bored", "ETSY_family", "ETSY_busy", "SITE_family", "KDP_plays", "CRS_single", "BB_site"):
    ur = UE_ROW[pid]
    put(wb_, f"A{r}", f"={q('Unit Economics')}!B{ur}&\" - \"&{q('Unit Economics')}!C{ur}")
    put(wb_, f"B{r}", f"={q('Unit Economics')}!M{ur}", fmt=CUR2)
    put(wb_, f"C{r}", f"=IF(B{r}<=0,\"-\",{fixed12}/B{r})", fmt=NUM0)
    put(wb_, f"D{r}", f"=IF(B{r}<=0,\"-\",C{r}/30)", fmt=NUM1)
    put(wb_, f"E{r}", f"Unit Economics row {ur}", font=F_NOTE)
    r += 1
r += 1
sec(wb_, r, 1, 5, "Retail wave self-funding check (Faire, Walmart Marketplace, Target Plus)"); r += 1
put(wb_, f"A{r}", "Retail monthly fixed costs once live (Retail Readiness Costs, section C)")
put(wb_, f"B{r}", f"={RR['C']}", fmt=CUR0); rC = r; r += 1
put(wb_, f"A{r}", "Blended net per retail unit (Unit Economics channel blend)")
put(wb_, f"B{r}", f"={UE['RETAIL']['net']}", fmt=CUR2); rN = r; r += 1
put(wb_, f"A{r}", "Retail units per month needed to cover retail fixed costs")
put(wb_, f"B{r}", f"=IF(B{rN}<=0,\"-\",B{rC}/B{rN})", fmt=NUM0); r += 1
put(wb_, f"A{r}", "One-time readiness cost (section B)")
put(wb_, f"B{r}", f"={RR['B']}", fmt=CUR0); rB = r; r += 1
put(wb_, f"A{r}", "Extra retail units to repay the one-time readiness cost")
put(wb_, f"B{r}", f"=IF(B{rN}<=0,\"-\",B{rB}/B{rN})", fmt=NUM0); r += 1
for pid in ("RT_faire", "RT_wmt", "RT_tgt"):
    ur = UE_ROW[pid]
    put(wb_, f"A{r}", f"=\"Net per unit: \"&{q('Unit Economics')}!C{ur}")
    put(wb_, f"B{r}", f"={q('Unit Economics')}!M{ur}", fmt=CUR2); r += 1
put(wb_, f"A{r}", "Wholesale at half of a $12.99 retail price leaves little after commission, landed cost and handling. A series, a 3-pack and a lower landed cost from a larger run change this.", font=F_NOTE)
wb_.column_dimensions["A"].width = 70
for c in "BCD":
    wb_.column_dimensions[c].width = 18
wb_.column_dimensions["E"].width = 60
wb_.freeze_panes = "B6"

# =====================================================================
# DASHBOARD
# =====================================================================
wd = ws_["Dashboard"]
put(wd, "A1", "Play Before Pixels: financial model dashboard", font=Font(name=AR, bold=True, size=16, color="1D2940"))
put(wd, "A2", "AlphaPlay LLC (Maryland) d/b/a Play Before Pixels. Draft for the founder, September 28, 2026. Planning estimates only: "
              "these are not forecasts of income or promises of results. Change any blue cell on Assumptions and every tab updates.", font=F_NOTE)
wd.merge_cells("A2:E3"); wd["A2"].alignment = WRAP; wd.row_dimensions[2].height = 30
hdr(wd, 5, 1, 5, ["Measure", "Conservative", "Expected", "Strong", "Where it comes from"], 22)
DR = {}
r = 6
def d_row(key, label, fn, fmt, note="", bold=False):
    global r
    put(wd, f"A{r}", label, bold=bold)
    for s in range(3):
        put(wd, f"{'BCD'[s]}{r}", fn(s), fmt=fmt, bold=bold)
    put(wd, f"E{r}", note, font=F_NOTE)
    DR[key] = r
    r += 1
def fy(s, key, k):
    return f"={q(FS)}!{YC[k]}{FR[s][key]}"
sec(wd, r, 1, 5, "Sales and profit by year (model year 1 = Oct 2026 - Sep 2027)"); r += 1
for yk in (1, 2, 3):
    d_row(f"gross{yk}", f"Gross sales, year {yk}", lambda s, yk=yk: fy(s, "gross", yk), CUR0, "36-Month Forecast" if yk == 1 else "")
for yk in (1, 2, 3):
    d_row(f"contrib{yk}", f"Net contribution, year {yk}", lambda s, yk=yk: fy(s, "contrib", yk), CUR0,
          "After platform, payment, print/POD and fulfilment costs" if yk == 1 else "")
for yk in (1, 2, 3):
    d_row(f"netop{yk}", f"Net operating result, year {yk}", lambda s, yk=yk: fy(s, "netop", yk), CUR0,
          "Contribution - ads - operating costs - wave fixed costs" if yk == 1 else "", bold=True)
d_row("oneoff", "One-time costs over 36 months (startup + wave launch + retail readiness)",
      lambda s: f"={q(FS)}!{YC['T']}{FR[s]['startup']}+{q(FS)}!{YC['T']}{FR[s]['waveone']}", CUR0, "Startup Costs; Retail Readiness Costs")
d_row("inv", "Inventory purchases over 36 months (board-book print runs)", lambda s: f"={q(FS)}!{YC['T']}{FR[s]['inv']}", CUR0, "Sized from pre-sale orders")
d_row("cum36", "Cumulative net result after one-time costs, month 36", lambda s: f"={q(FS)}!{mc(36)}{FR[s]['cumnet']}", CUR0, "", bold=True)
r += 1
sec(wd, r, 1, 5, "Cash and break-even"); r += 1
d_row("peak", "Peak funding need (lowest operating balance)", lambda s: f"={q('Break-even')}!{'BCD'[s]}{BE['peak']}", CUR0, "Cash Flow (includes payout delays)", bold=True)
d_row("peakm", "   ...in month", lambda s: dt(f"{q('Break-even')}!{'BCD'[s]}{BE['peak_m']}"), DATEF)
d_row("be", "Sustained operating break-even", lambda s: f"={q('Break-even')}!{'BCD'[s]}{BE['sust_d']}", DATEF, "Break-even tab", bold=True)
d_row("pay", "Cash payback (operating balance stays >= 0)", lambda s: f"={q('Break-even')}!{'BCD'[s]}{BE['cash_d']}", DATEF)
d_row("cash36", "Total cash at month 36 (operating + tax reserve)", lambda s: f"={q('Cash Flow')}!{mc(36)}{CFR[s]['total']}", CUR0)
d_row("gap36", "Operating balance vs 3-month reserve target, month 36", lambda s: f"={q('Cash Flow')}!{mc(36)}{CFR[s]['gap']}", CUR0, "ops/ROUTINE.md reserve target")
r += 1
sec(wd, r, 1, 5, "Volume and mix"); r += 1
d_row("orders36", "Orders in month 36", lambda s: f"={q(FS)}!{mc(36)}{FR[s]['orders']}", NUM0)
d_row("aov", "Average gross sale per order, year 3", lambda s: f"=IF({q(FS)}!{YC[3]}{FR[s]['orders']}=0,0,{q(FS)}!{YC[3]}{FR[s]['gross']}/{q(FS)}!{YC[3]}{FR[s]['orders']})", CUR2)
d_row("list36", "Email subscribers at month 36", lambda s: f"={q(FS)}!{mc(36)}{FR[s]['email_end']}", NUM0)
d_row("digshare", "Share of 36-month contribution from digital (site, Etsy, MoR, course)",
      lambda s: f"=IF({q(FS)}!{YC['T']}{FR[s]['contrib']}=0,0,({q(FS)}!{YC['T']}{FR[s]['SITE_contrib']}+{q(FS)}!{YC['T']}{FR[s]['ETSY_contrib']}+{q(FS)}!{YC['T']}{FR[s]['MOR_contrib']}+{q(FS)}!{YC['T']}{FR[s]['COURSE_contrib']})/{q(FS)}!{YC['T']}{FR[s]['contrib']})", PCT1)
d_row("bookshare", "Share from print-on-demand books (KDP, IngramSpark)",
      lambda s: f"=IF({q(FS)}!{YC['T']}{FR[s]['contrib']}=0,0,({q(FS)}!{YC['T']}{FR[s]['KDP_contrib']}+{q(FS)}!{YC['T']}{FR[s]['INGRAM_contrib']})/{q(FS)}!{YC['T']}{FR[s]['contrib']})", PCT1)
d_row("bbshare", "Share from the board book (Wave 3)",
      lambda s: f"=IF({q(FS)}!{YC['T']}{FR[s]['contrib']}=0,0,{q(FS)}!{YC['T']}{FR[s]['BOARD_contrib']}/{q(FS)}!{YC['T']}{FR[s]['contrib']})", PCT1)
d_row("schshare", "Share from the school/group wave (held for counsel)",
      lambda s: f"=IF({q(FS)}!{YC['T']}{FR[s]['contrib']}=0,0,{q(FS)}!{YC['T']}{FR[s]['SCHOOL_contrib']}/{q(FS)}!{YC['T']}{FR[s]['contrib']})", PCT1)
d_row("rtshare", "Share from retail and wholesale (Faire, Walmart, Target Plus)",
      lambda s: f"=IF({q(FS)}!{YC['T']}{FR[s]['contrib']}=0,0,{q(FS)}!{YC['T']}{FR[s]['RETAIL_contrib']}/{q(FS)}!{YC['T']}{FR[s]['contrib']})", PCT1)
d_row("presale", "Board-book pre-sale copies before the print month", lambda s: f"={q(FS)}!$B${FR[s]['presale']}", NUM0, "Wave 3 pre-sale (Feb-Apr 2027)")
d_row("need", "   ...copies needed to fund the first run (go / no-go line)", lambda s: f"={q(FS)}!$B${FR[s]['need']}", NUM0, "Below this line, delay the print run and keep the POD paperback")
d_row("onhand", "Board-book copies at the 3PL, month 36", lambda s: f"={q(FS)}!{mc(36)}{FR[s]['onhand']}", NUM0, "Cash tied up in stock")
r += 1
put(wd, f"A{r}", "How to use", bold=True); r += 1
for t in ["1. Edit blue cells only. Yellow cells on Assumptions are the levers that move results most (traffic, conversion, include flags, IngramSpark discount, owner contribution).",
          "2. The cost position on Assumptions (0 = every cost at the low end of its range, 1 = high end) tests cost risk in one step.",
          "3. School-facing products stay off in Conservative and start only when employment counsel clears them. Retail (including Target Plus) starts only after a 12-month record.",
          "4. Every [VERIFY] figure must be checked on the live page or a written quote before money is committed.",
          "5. Replace the traffic and conversion guesses with the first 90 days of real data from the Friday scorecard."]:
    put(wd, f"A{r}", t, font=F_NOTE); r += 1
for c, w in zip("ABCDE", [62, 16, 16, 16, 44]):
    wd.column_dimensions[c].width = w
wd.freeze_panes = "B6"

# Charts
def line_chart(title_, rows_fn, sheet, ytitle, anchor):
    ch = LineChart()
    ch.title = title_
    ch.y_axis.title = ytitle
    ch.y_axis.number_format = '$#,##0'
    ch.x_axis.number_format = 'mmm-yy'
    ch.height, ch.width = 8, 18
    ws = ws_[sheet]
    cats = Reference(ws, min_col=FC, max_col=FC + 35, min_row=5, max_row=5)
    for s in range(3):
        row = rows_fn(s)
        ref = Reference(ws, min_col=FC, max_col=FC + 35, min_row=row, max_row=row)
        se = Series(ref, title=SCN[s])
        se.graphicalProperties.line.solidFill = SERIES_COLORS[s]
        se.graphicalProperties.line.width = 25400
        se.smooth = False
        ch.series.append(se)
    ch.set_categories(cats)
    ch.x_axis.tickLblPos = "low"
    ch.legend.position = "b"
    wd.add_chart(ch, anchor)
line_chart("Monthly gross sales by scenario", lambda s: FR[s]["gross"], FS, "$ per month", "G5")
line_chart("Operating account balance by scenario (month end)", lambda s: CFR[s]["opbal"], "Cash Flow", "$", "G22")
r += 1
put(wd, f"A{r}", "Chart data: net operating result by model year", bold=True); r += 1
cd0 = r
for yk in (1, 2, 3):
    put(wd, f"A{r}", f"Year {yk}")
    for s in range(3):
        put(wd, f"{'BCD'[s]}{r}", f"={'BCD'[s]}{DR[f'netop{yk}']}", fmt=CUR0)
    r += 1
bc = BarChart()
bc.type = "col"; bc.grouping = "clustered"
bc.title = "Net operating result by model year"
bc.y_axis.number_format = '$#,##0'
bc.height, bc.width = 8, 18
for s in range(3):
    ref = Reference(wd, min_col=2 + s, max_col=2 + s, min_row=cd0, max_row=cd0 + 2)
    se = Series(ref, title=SCN[s])
    se.graphicalProperties.solidFill = SERIES_COLORS[s]
    se.graphicalProperties.line.solidFill = SERIES_COLORS[s]
    bc.series.append(se)
bc.set_categories(Reference(wd, min_col=1, max_col=1, min_row=cd0, max_row=cd0 + 2))
bc.x_axis.tickLblPos = "low"
bc.legend.position = "b"
wd.add_chart(bc, "G39")

from openpyxl.worksheet.properties import PageSetupProperties
for n in SHEETS:
    w = ws_[n]
    w.page_setup.orientation = "landscape"
    w.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
# tab colors
for n in SHEETS:
    ws_[n].sheet_properties.tabColor = "1D2940" if n in ("Dashboard", "Assumptions") else "8A94A8"
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
